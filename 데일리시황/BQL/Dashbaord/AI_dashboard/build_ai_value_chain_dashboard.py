"""Build the local AI value-chain price / 12MF EPS / 12MF P/E dashboard.

The classification is imported from the maintained AI report builder.  The
source workbook is read-only; this script does not alter BQuant_Master.xlsx.
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

import numpy as np
import openpyxl


MIN_CAP_COVERAGE = 0.60
MAX_STOCK_PE = 200.0


def positive_finite(value: np.ndarray) -> np.ndarray:
    return np.isfinite(value) & (value > 0)


def as_json_number(value: float) -> float | None:
    return round(float(value), 5) if math.isfinite(float(value)) else None


def first_last(values: list[float | None], dates: list[str]) -> tuple[str | None, str | None]:
    positions = [i for i, value in enumerate(values) if value is not None]
    if not positions:
        return None, None
    return dates[positions[0]], dates[positions[-1]]


def rebase(values: np.ndarray) -> list[float | None]:
    output: list[float | None] = [None] * len(values)
    valid = np.flatnonzero(np.isfinite(values) & (values > 0))
    if not len(valid):
        return output
    base = float(values[valid[0]])
    for i in valid:
        output[int(i)] = as_json_number(float(values[i]) / base * 100.0)
    return output


def summarize_group(
    rows: list[int],
    dates: list[str],
    prices: np.ndarray,
    caps: np.ndarray,
    eps: np.ndarray,
    pes: np.ndarray,
) -> dict:
    n = len(rows)
    if not n:
        empty = [None] * len(dates)
        return {"members": 0, "price": empty, "eps": empty, "pe": empty,
                "coverage": {"price": empty, "eps": empty, "pe": empty},
                "range": {"price": [None, None], "eps": [None, None], "pe": [None, None]}}

    p = prices[rows, :]
    c = caps[rows, :]
    e = eps[rows, :]
    pe = pes[rows, :]
    base = positive_finite(p) & positive_finite(c)
    total_base_cap = np.sum(np.where(base, c, 0.0), axis=0)
    shares = np.divide(c, p, out=np.zeros_like(c), where=base)

    def coverage(mask: np.ndarray) -> np.ndarray:
        participating = np.sum(np.where(mask, c, 0.0), axis=0)
        return np.divide(participating, total_base_cap,
                         out=np.zeros_like(participating), where=total_base_cap > 0)

    price_count = np.sum(base, axis=0)
    price_cov = np.divide(price_count, n, dtype=float)
    aggregate_price = np.divide(
        total_base_cap,
        np.sum(shares, axis=0),
        out=np.full(len(dates), np.nan),
        where=np.sum(shares, axis=0) > 0,
    )
    aggregate_price[price_cov < MIN_CAP_COVERAGE] = np.nan

    eps_valid = base & np.isfinite(e) & (np.abs(e) < 100_000)
    eps_cov = coverage(eps_valid)
    total_earnings = np.sum(np.where(eps_valid, e * shares, 0.0), axis=0)
    eps_shares = np.sum(np.where(eps_valid, shares, 0.0), axis=0)
    basket_eps = np.divide(total_earnings, eps_shares,
                           out=np.full(len(dates), np.nan), where=eps_shares > 0)
    basket_eps[(eps_cov < MIN_CAP_COVERAGE) | (basket_eps <= 0)] = np.nan

    pe_valid = base & np.isfinite(pe) & (pe >= 1.0) & (pe <= MAX_STOCK_PE)
    pe_cov = coverage(pe_valid)
    pe_cap = np.sum(np.where(pe_valid, c, 0.0), axis=0)
    earnings_yield_cap = np.sum(
        np.divide(c, pe, out=np.zeros_like(c), where=pe_valid), axis=0
    )
    harmonic_pe = np.divide(pe_cap, earnings_yield_cap,
                            out=np.full(len(dates), np.nan), where=earnings_yield_cap > 0)
    harmonic_pe[pe_cov < MIN_CAP_COVERAGE] = np.nan

    result_price = rebase(aggregate_price)
    result_eps = rebase(basket_eps)
    result_pe = [as_json_number(x) for x in harmonic_pe]
    cover = {
        "price": [as_json_number(x) for x in price_cov],
        "eps": [as_json_number(x) for x in eps_cov],
        "pe": [as_json_number(x) for x in pe_cov],
    }
    return {
        "members": n,
        "price": result_price,
        "eps": result_eps,
        "pe": result_pe,
        "coverage": cover,
        "range": {k: list(first_last(v, dates)) for k, v in
                  (("price", result_price), ("eps", result_eps), ("pe", result_pe))},
    }


def build_payload(raw, stages, stock_symbol, ai_builder) -> dict:
    price_all = raw.metrics["price"].copy()
    cap_all = raw.metrics["mcap"].copy()
    eps_all = raw.metrics["eps12"].copy()
    pe_all = raw.metrics["pe"].copy()
    ai_builder.reconcile_known_splits(price_all, raw.tickers)
    latest = ai_builder.latest_trading_index(price_all)
    sessions = [i for i in ai_builder.actual_session_indices(price_all) if i <= latest]
    if len(sessions) < 2:
        raise RuntimeError("SPX Price에 유효한 거래일이 두 개 미만입니다.")
    dates = [raw.dates[i].isoformat() for i in sessions]
    price = price_all[:, sessions]
    cap = cap_all[:, sessions]
    eps = eps_all[:, sessions]
    pe = pe_all[:, sessions]
    symbol_to_row = {stock_symbol(ticker): i for i, ticker in enumerate(raw.tickers)}
    groups = []
    for stage in stages:
        mode_data = {}
        missing = {}
        for mode, symbols in (
            ("Core", stage["core"]),
            ("Expanded", stage["core"] + stage["adjacent"]),
        ):
            rows = [symbol_to_row[symbol] for symbol in symbols if symbol in symbol_to_row]
            missing[mode] = [symbol for symbol in symbols if symbol not in symbol_to_row]
            mode_data[mode] = summarize_group(rows, dates, price, cap, eps, pe)
        groups.append({
            "key": stage["key"], "name": stage["name"], "short": stage["short"],
            "color": stage["color"], "description": stage["description"],
            "coreTickers": stage["core"], "adjacentTickers": stage["adjacent"],
            "missing": missing, "modes": mode_data,
        })
    return {
        "meta": {
            "source": str(raw.source), "sourceSheet": raw.sheet,
            "index": raw.code, "currency": raw.currency,
            "asOf": dates[-1], "firstSession": dates[0],
            "lastSession": dates[-1], "sessions": len(dates),
            "minCapCoverage": MIN_CAP_COVERAGE,
            "method": "Price=shares-weighted basket price, EPS=shares-weighted 12MF blended EPS rebased to 100, PE=harmonic market-cap weighted 12MF PE",
        },
        "dates": dates,
        "groups": groups,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bql-root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--master", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = args.bql_root.resolve()
    master = (args.master or root / "Rawfile" / "BQuant_Master.xlsx").resolve()
    output = (args.output or root / "Dashbaord" / "AI_dashboard" / "AI_Value_Chain_Dashboard.html").resolve()
    if not master.is_file():
        raise FileNotFoundError(master)
    sys.path.insert(0, str(root / "Theme" / "Material" / "AI"))
    sys.path.insert(0, str(root / "Dashbaord"))
    import build_sp500_ai_value_chain as ai_builder  # noqa: E402
    workbook = openpyxl.load_workbook(master, read_only=True, data_only=True)
    try:
        sheet_name = next((name for name in workbook.sheetnames if name.casefold() == "spx raw"), None)
        if sheet_name is None:
            raise RuntimeError("마스터 파일에 SPX Raw 시트가 없습니다.")
        rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()
    raw = ai_builder.raw_builder.parse_raw_rows(rows, master.name, sheet_name)
    if raw is None or raw.code != "SPX":
        raise RuntimeError("SPX Raw 시트를 파싱하지 못했습니다.")
    payload = build_payload(raw, ai_builder.STAGES, ai_builder.ticker_symbol, ai_builder)
    template = (Path(__file__).resolve().parent / "dashboard_template.html").read_text(encoding="utf-8")
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    serialized = serialized.replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    html = template.replace("/*__DASHBOARD_DATA__*/", serialized)
    if html == template:
        raise RuntimeError("dashboard_template.html에 데이터 삽입 위치가 없습니다.")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html, encoding="utf-8")
    print(json.dumps({"output": str(output), "asOf": payload["meta"]["asOf"],
                      "firstSession": payload["meta"]["firstSession"],
                      "sessions": len(payload["dates"]), "groups": len(payload["groups"]),
                      "bytes": output.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
