from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\infomax\Documents\ChatGPT\컨텍스트매크로")
OUT = ROOT / "artifacts" / "europe_topix_ai_comparison"
OUTPUT_DOCX = OUT / "유럽_TOPIX_산업재_AI_Capital_Goods_비교_20260904.docx"
EUROPE_JSON = Path(
    r"C:\Users\infomax\Documents\python\QAE\데일리시황\BQL\Theme\Sector_Industry\Reports\output\data\STOXX600_TOPIX_Sector_Industry_5D_20260904.json"
)
JAPAN_JSON = ROOT / "artifacts" / "topix_industrials_ai" / (
    "topix_capital_goods_basket_analysis_20260828_20260904.json"
)
EUROPE_MEMO_DIR = ROOT / "artifacts" / "stoxx_industrials_ai"

NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
RED = "C00000"
GREEN = "008000"
AMBER = "9C6500"
GRAY = "666666"
LIGHT = "D9E2F3"
ALT = "F6F9FC"


EUROPE_DEFENSE = {
    "CSG NA Equity", "RHM GY Equity", "R3NK GY Equity", "HAG GY Equity",
    "SAABB SS Equity", "BAB LN Equity", "QQ/ LN Equity", "BA/ LN Equity",
    "LDO IM Equity", "HO FP Equity", "AM FP Equity", "KOG NO Equity",
}
EUROPE_AI = {
    "SU FP Equity", "ABBN SE Equity", "SIE GY Equity", "LR FP Equity",
    "PRY IM Equity", "BEAN SE Equity", "ALFA SS Equity", "IMI LN Equity",
    "VACN SE Equity", "ATCOA SS Equity", "HUBN SE Equity", "ENR GY Equity",
    "NKT DC Equity", "NEX FP Equity", "RXL FP Equity",
}


def install_font() -> None:
    for candidate in [Path(r"C:\Windows\Fonts\malgun.ttf"), Path(r"C:\Windows\Fonts\arial.ttf")]:
        if candidate.exists():
            font_manager.fontManager.addfont(str(candidate))
            plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(candidate)).get_name()
            break
    plt.rcParams["axes.unicode_minus"] = False


def group_stat(rows: list[dict], total_cap: float) -> dict[str, float | int]:
    values = np.asarray([float(row["return5d"]) for row in rows], dtype=float)
    caps = np.asarray([float(row["startMarketCap"]) for row in rows], dtype=float)
    return {
        "n": len(rows),
        "ew": float(values.mean()),
        "median": float(np.median(values)),
        "breadth": float((values > 0).mean()),
        "cw": float(np.average(values, weights=caps)),
        "cap_share": float(caps.sum() / total_cap),
    }


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:+.{digits}f}%"


def pp(value: float, digits: int = 2) -> str:
    return f"{value * 100:+.{digits}f}%p"


def clean_name(value: str) -> str:
    for suffix in [" Corp", " Co Ltd", " Ltd", " Inc", " Holdings", " KK"]:
        value = value.replace(suffix, "")
    return value.strip()


def load_data() -> tuple[dict, dict, dict, list[dict]]:
    europe_raw = json.loads(EUROPE_JSON.read_text(encoding="utf-8"))
    europe_market = next(row for row in europe_raw["markets"] if row["key"] == "STOXX600")
    capital_goods = [
        row for row in europe_market["stocks"]
        if row["sector"] == "Industrials" and row["industry"] == "Capital Goods" and not row["excluded"]
    ]
    total_cap = sum(float(row["startMarketCap"]) for row in capital_goods)
    europe = {
        "All": group_stat(capital_goods, total_cap),
        "Defense": group_stat([row for row in capital_goods if row["ticker"] in EUROPE_DEFENSE], total_cap),
        "AI_Combined": group_stat([row for row in capital_goods if row["ticker"] in EUROPE_AI], total_cap),
        "Other": group_stat(
            [row for row in capital_goods if row["ticker"] not in EUROPE_DEFENSE | EUROPE_AI], total_cap
        ),
    }
    japan_raw = json.loads(JAPAN_JSON.read_text(encoding="utf-8"))
    return europe, japan_raw["groupSummary"], japan_raw["meta"], japan_raw["stocks"]


def make_charts(europe: dict, japan: dict, japan_stocks: list[dict]) -> dict[str, Path]:
    OUT.mkdir(parents=True, exist_ok=True)
    install_font()
    chart_paths: dict[str, Path] = {}

    keys = ["All", "Defense", "AI_Combined", "Other"]
    labels = ["Capital Goods 전체", "방산", "AI·데이터센터 광의", "기타"]
    eu = [float(europe[key]["ew"]) * 100 for key in keys]
    jp = [float(japan[key]["equalWeight"]) * 100 for key in keys]
    fig, ax = plt.subplots(figsize=(11.4, 5.8), dpi=190)
    y = np.arange(len(labels))
    h = 0.34
    ax.barh(y - h / 2, eu, height=h, color="#17365D", label="STOXX Europe 600")
    ax.barh(y + h / 2, jp, height=h, color="#2F75B5", label="TOPIX")
    ax.axvline(0, color="#555555", linewidth=0.8)
    for pos, value in zip(y - h / 2, eu):
        ax.text(value - 0.12, pos, f"{value:.1f}%", va="center", ha="right", color="white", fontsize=10, fontweight="bold")
    for pos, value in zip(y + h / 2, jp):
        ax.text(value - 0.12, pos, f"{value:.1f}%", va="center", ha="right", color="white", fontsize=10, fontweight="bold")
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlim(min(eu + jp) - 0.9, 0.7)
    ax.grid(axis="x", color="#D9E2F3", linewidth=0.8)
    ax.spines[["top", "right", "left", "bottom"]].set_visible(False)
    ax.legend(loc="lower right", frameon=False, ncol=2)
    fig.suptitle("유럽과 일본 Capital Goods: 동일가중 5일 수익률", x=0.06, y=0.98, ha="left", fontsize=17, fontweight="bold", color="#17365D")
    fig.text(0.06, 0.925, "2026-08-28~2026-09-04 | 현지통화 종가 | 방산·AI·기타는 상호배타적 분석 프록시", fontsize=9, color="#666666")
    fig.tight_layout(rect=(0, 0, 1, 0.89))
    chart_paths["cross"] = OUT / "01_cross_market_group_returns.png"
    fig.savefig(chart_paths["cross"], bbox_inches="tight", facecolor="white")
    plt.close(fig)

    subkeys = [
        "Defense", "Direct_Cable", "Direct_Power_Cooling_Equipment",
        "Adjacent_Semicap_FA", "Direct_Construction_HVAC", "Adjacent_Grid_Electrical",
    ]
    sublabels = ["방산", "광·케이블", "전력·냉각장비", "반도체장비·FA", "시공·HVAC", "전력망·전기공사"]
    ew = [float(japan[key]["equalWeight"]) * 100 for key in subkeys]
    cw = [float(japan[key]["capWeight"]) * 100 for key in subkeys]
    fig, ax = plt.subplots(figsize=(11.4, 6.5), dpi=190)
    y = np.arange(len(sublabels))
    ax.barh(y - h / 2, ew, height=h, color="#2F75B5", label="동일가중")
    ax.barh(y + h / 2, cw, height=h, color="#95B3D7", label="시총가중 진단")
    ax.axvline(0, color="#555555", linewidth=0.8)
    for pos, value in zip(y - h / 2, ew):
        ax.text(value + (0.10 if value >= 0 else -0.10), pos, f"{value:+.1f}%", va="center", ha="left" if value >= 0 else "right", color="#17365D", fontsize=9, fontweight="bold")
    for pos, value in zip(y + h / 2, cw):
        ax.text(value + (0.10 if value >= 0 else -0.10), pos, f"{value:+.1f}%", va="center", ha="left" if value >= 0 else "right", color="#17365D", fontsize=9, fontweight="bold")
    ax.set_yticks(y, sublabels)
    ax.invert_yaxis()
    ax.set_xlim(min(ew + cw) - 0.9, max(ew + cw) + 0.9)
    ax.grid(axis="x", color="#D9E2F3", linewidth=0.8)
    ax.spines[["top", "right", "left", "bottom"]].set_visible(False)
    ax.legend(loc="lower right", frameon=False, ncol=2)
    fig.suptitle("TOPIX Capital Goods 내부: AI 전반이 아니라 장비군 약세", x=0.06, y=0.98, ha="left", fontsize=17, fontweight="bold", color="#17365D")
    fig.text(0.06, 0.93, "광·케이블, 전력·냉각장비, 반도체장비·FA는 약세 지속; 국내 전력망·전기공사는 플러스", fontsize=9, color="#666666")
    fig.tight_layout(rect=(0, 0, 1, 0.90))
    chart_paths["subthemes"] = OUT / "02_topix_ai_subthemes.png"
    fig.savefig(chart_paths["subthemes"], bbox_inches="tight", facecolor="white")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(11.4, 5.9), dpi=190)
    for key, label, color in [
        ("Defense", "방산", "#C00000"),
        ("AI_Combined", "AI·데이터센터 광의", "#2F75B5"),
        ("Other", "기타 자본재", "#7F7F7F"),
    ]:
        series = japan[key]["path"]
        dates = [row["date"][5:].replace("-", "/") for row in series]
        values = [float(row["ewCumulative"]) * 100 for row in series]
        ax.plot(dates, values, marker="o", linewidth=2.5, label=label, color=color)
    ax.axhline(0, color="#555555", linewidth=0.8)
    ax.axvspan(3 - 0.35, 3 + 0.35, color="#FCE4D6", alpha=0.75)
    ax.text(3, -7.0, "9/2 BOJ 매파 발언\nJGB 금리 급등", ha="center", va="bottom", color="#9C0006", fontsize=9, fontweight="bold")
    ax.set_ylim(-7.8, 2.2)
    ax.set_ylabel("8/28 대비 누적수익률 (%)")
    ax.grid(color="#D9E2F3", linewidth=0.8)
    ax.spines[["top", "right", "left", "bottom"]].set_visible(False)
    ax.legend(loc="lower left", frameon=False, ncol=3)
    fig.suptitle("TOPIX Capital Goods 하락은 9월 2일에 집중", x=0.06, y=0.98, ha="left", fontsize=17, fontweight="bold", color="#17365D")
    fig.text(0.06, 0.925, "종목별 8/28 종가=0, 동일가중 누적수익률 | 산업생산 발표(8/31)는 악재가 아니라 예상 상회", fontsize=9, color="#666666")
    fig.tight_layout(rect=(0, 0, 1, 0.89))
    chart_paths["path"] = OUT / "03_topix_daily_path.png"
    fig.savefig(chart_paths["path"], bbox_inches="tight", facecolor="white")
    plt.close(fig)

    downside = {
        "STOXX Europe 600": [26.3, 43.1, 30.6],
        "TOPIX": [
            float(japan["Defense"]["shareOfGrossDownside"]) * 100,
            float(japan["AI_Combined"]["shareOfGrossDownside"]) * 100,
            float(japan["Other"]["shareOfGrossDownside"]) * 100,
        ],
    }
    fig, ax = plt.subplots(figsize=(11.4, 3.6), dpi=190)
    colors = ["#C00000", "#2F75B5", "#A6A6A6"]
    labels_stack = ["방산", "AI·DC 광의", "기타"]
    for y_pos, (market, values) in enumerate(downside.items()):
        left = 0.0
        for label, value, color in zip(labels_stack, values, colors):
            ax.barh(y_pos, value, left=left, color=color, height=0.48, label=label if y_pos == 0 else None)
            if value >= 10:
                ax.text(left + value / 2, y_pos, f"{value:.1f}%", ha="center", va="center", color="white", fontsize=10, fontweight="bold")
            left += value
    ax.set_yticks([0, 1], list(downside.keys()))
    ax.invert_yaxis()
    ax.set_xlim(0, 100)
    ax.set_xlabel("Capital Goods gross downside 기여 비중")
    ax.spines[["top", "right", "left", "bottom"]].set_visible(False)
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.7)
    ax.legend(loc="lower center", bbox_to_anchor=(0.5, -0.42), frameon=False, ncol=3)
    fig.suptitle("AI 관련 대형주가 포트폴리오 손실의 큰 축", x=0.06, y=1.02, ha="left", fontsize=16, fontweight="bold", color="#17365D")
    fig.tight_layout(rect=(0, 0.05, 1, 0.92))
    chart_paths["downside"] = OUT / "04_cross_market_downside_share.png"
    fig.savefig(chart_paths["downside"], bbox_inches="tight", facecolor="white")
    plt.close(fig)

    contributors = sorted(
        [row for row in japan_stocks if float(row.get("downsideContributionPp") or 0) < 0],
        key=lambda row: float(row["downsideContributionPp"]),
    )[:12]
    contributors = list(reversed(contributors))
    names = [clean_name(row["name"]) for row in contributors]
    values = [float(row["downsideContributionPp"]) * 100 for row in contributors]
    colors = ["#C00000" if row["bucket"] == "Defense" else "#2F75B5" if row["bucket"].startswith("AI") else "#7F7F7F" for row in contributors]
    fig, ax = plt.subplots(figsize=(11.4, 7.0), dpi=190)
    bars = ax.barh(np.arange(len(names)), values, color=colors, height=0.62)
    ax.axvline(0, color="#555555", linewidth=0.8)
    ax.set_yticks(np.arange(len(names)), names)
    for bar, value in zip(bars, values):
        ax.text(value - 0.008, bar.get_y() + bar.get_height() / 2, f"{value:.2f}%p", va="center", ha="right", color="white", fontsize=9, fontweight="bold")
    ax.grid(axis="x", color="#D9E2F3", linewidth=0.8)
    ax.spines[["top", "right", "left", "bottom"]].set_visible(False)
    fig.suptitle("TOPIX Capital Goods 시총가중 하락 기여 상위", x=0.06, y=0.98, ha="left", fontsize=17, fontweight="bold", color="#17365D")
    fig.text(0.06, 0.93, "9/2 시가총액으로 8/28 시가총액을 역산한 진단치 | 파랑 AI·DC, 빨강 방산, 회색 기타", fontsize=9, color="#666666")
    fig.tight_layout(rect=(0, 0, 1, 0.90))
    chart_paths["contributors"] = OUT / "05_topix_downside_contributors.png"
    fig.savefig(chart_paths["contributors"], bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return chart_paths


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_run(run, size: float = 9.3, bold: bool = False, color: str = "111111") -> None:
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text: str, url: str, color: str = "0563C1") -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Malgun Gothic")
    r_fonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_paragraph(doc: Document, text: str = "", *, size: float = 9.3, color: str = "111111", bold: bool = False, bullet: bool = False, after: float = 4, keep: bool = False):
    style = "List Bullet" if bullet else None
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    style_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_tagged(doc: Document, tag: str, text: str, *, after: float = 4) -> None:
    color = {"FACT": BLUE, "INFERENCE": AMBER, "SPECULATION": RED}[tag]
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    style_run(p.add_run(f"[{tag}] "), size=9.2, bold=True, color=color)
    style_run(p.add_run(text), size=9.2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_callout(doc: Document, label: str, text: str, fill: str = PALE, accent: str = BLUE) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    set_cell_width(cell, 18.1)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Cm(0.12)
    style_run(p.add_run(f"{label}  "), size=9.8, bold=True, color=accent)
    style_run(p.add_run(text), size=9.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: list[str], rows: Iterable[list[str]], widths: list[float], *, font_size: float = 7.8) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_header(header)
    prevent_split(header)
    for idx, (text, width) in enumerate(zip(headers, widths)):
        cell = header.cells[idx]
        set_cell_width(cell, width)
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(text), size=font_size, bold=True, color="FFFFFF")
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_split(row)
        for col_idx, (text, width) in enumerate(zip(values, widths)):
            cell = row.cells[col_idx]
            set_cell_width(cell, width)
            if row_idx % 2:
                set_cell_shading(cell, ALT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            color = RED if isinstance(text, str) and text.startswith("-") else GREEN if isinstance(text, str) and text.startswith("+") else "111111"
            style_run(p.add_run(str(text)), size=font_size, bold=(col_idx == 0), color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc: Document, path: Path, width: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    style_run(cap.add_run(caption), size=7.2, color=GRAY)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(paragraph.add_run("2026-09-08  |  "), size=7.2, color=GRAY)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def source_line(doc: Document, title: str, url: str, note: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    add_hyperlink(p, title, url)
    style_run(p.add_run(f" — {note}"), size=7.8, color=GRAY)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(9.3)
    for style_name, size, color in [("Title", 23, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 11.5, BLUE)]:
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("유럽·TOPIX · Industrials / Capital Goods · 5D"), size=7.2, color=GRAY)
    add_page_number(section.footer.paragraphs[0])


def names_by_subbucket(stocks: list[dict], key: str) -> str:
    selected = sorted([row for row in stocks if row["subbucket"] == key], key=lambda row: row["ticker"])
    return ", ".join(f"{clean_name(row['name'])} ({row['ticker'].split()[0]})" for row in selected)


def build_document(europe: dict, japan: dict, japan_meta: dict, japan_stocks: list[dict], charts: dict[str, Path]) -> None:
    doc = Document()
    configure_document(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(8)
    style_run(p.add_run("유럽·일본 Capital Goods 5D 비교"), size=23, bold=True, color=NAVY)
    add_paragraph(doc, "AI 약세는 있었나, 산업생산 부진이 원인이었나", size=14.5, bold=True, color=BLUE, after=7)
    add_paragraph(doc, "분석 구간 2026-08-28~2026-09-04  |  작성 2026-09-08", size=8.5, color=GRAY, after=10)
    add_callout(
        doc,
        "한 줄 결론",
        "두 시장 모두 AI 관련 Capital Goods가 기타 자본재보다 약했다. 그러나 산업생산 악화가 직접 촉매였다는 증거는 약하다. 유럽의 부진한 7월 산업생산은 가격 구간 뒤에 발표됐고, 일본의 7월 산업생산은 오히려 예상보다 좋았다. 공통된 1차 설명은 유가·물가·장기금리 상승에 따른 장기 성장주의 할인율 충격이다.",
    )
    add_tagged(doc, "FACT", "유럽 AI·데이터센터 광의 15개는 -2.39%, 일본 AI 광의 63개는 -2.05%였다. 각 시장의 기타 Capital Goods는 각각 -1.18%, -0.88%였다.")
    add_tagged(doc, "FACT", "일본에서는 방산 7개가 -4.89%로 종목당 낙폭은 더 컸지만, gross downside 기여는 AI 광의 61.6%, 방산 23.9%였다.")
    add_tagged(doc, "INFERENCE", "‘방산만 나빴다’도, ‘산업생산이 나빠 AI 장비주가 무너졌다’도 불완전하다. 약세는 방산과 AI 장비군에 함께 나타났지만 촉매는 실물지표보다 금리·밸류에이션에 더 가깝다.")
    add_image(doc, charts["cross"], 7.05, "그림 1. BQL·검증된 공개 종가 기준 동일가중 수익률. 두 시장 모두 현지통화 가격수익률이며 배당·환율은 제외.")

    doc.add_page_break()
    add_heading(doc, "Executive Summary")
    rows = []
    for label, key in [("Capital Goods 전체", "All"), ("방산", "Defense"), ("AI·DC 광의", "AI_Combined"), ("기타", "Other")]:
        rows.append([
            label,
            str(europe[key]["n"]), pct(float(europe[key]["ew"])), pct(float(europe[key]["cw"])),
            str(japan[key]["n"]), pct(float(japan[key]["equalWeight"])), pct(float(japan[key]["capWeight"])),
        ])
    add_table(doc, ["묶음", "EU N", "EU EW", "EU CW*", "JP N", "JP EW", "JP CW*"], rows, [4.3, 1.3, 2.0, 2.0, 1.3, 2.0, 2.0], font_size=7.6)
    add_paragraph(doc, "*시총가중은 공식 지수 비중이 아니라 시작 시가총액을 이용한 노출 진단값이다.", size=7.2, color=GRAY)
    add_tagged(doc, "FACT", f"AI 광의는 유럽에서 기타 대비 {pp(float(europe['AI_Combined']['ew']) - float(europe['Other']['ew']))}, 일본에서 {pp(float(japan['AI_Combined']['equalWeight']) - float(japan['Other']['equalWeight']))} 언더퍼폼했다.")
    add_tagged(doc, "FACT", "유럽 gross downside는 AI 43.1%, 방산 26.3%; 일본은 AI 61.6%, 방산 23.9%였다. 대형 AI 전력·자동화·냉각·케이블주가 시총가중 손실을 크게 만들었다.")
    add_image(doc, charts["downside"], 6.9, "그림 2. 각 그룹의 음(-)의 시총가중 기여를 합산한 gross downside 비중. 그룹 내 플러스 종목의 상쇄효과는 제외.")
    add_callout(doc, "PM 판독", "꼬리위험의 모양은 방산이 만들었지만, 포트폴리오 손익의 중심은 AI 관련 대형주였다. 일본에서는 이 구도가 유럽보다 더 선명했다.", fill="FFF2CC", accent=AMBER)

    doc.add_page_break()
    add_heading(doc, "TOPIX 날짜 오류: 사용자가 준 9월 4일 데이터는 실제로 있었다")
    add_table(
        doc,
        ["항목", "기존 산출물", "재검토 결과"],
        [
            ["표시 기준일", "2026-08-21", "2026-09-04"],
            ["사용 구간", "8/14~8/21", "8/28~9/4"],
            ["구성종목", "옛 fallback 1,636", "현재 Master 1,635"],
            ["원천 선택", "historical_fallback", "현재 구성 + 검증된 종가 복원"],
            ["현지 데이터", "9/2·9/3·9/4를 버림", "전부 보존·대조"],
        ],
        [3.2, 6.1, 7.3],
        font_size=8.0,
    )
    add_tagged(doc, "FACT", "Master에는 9월 4일까지의 TOPIX 데이터가 있었다. 다만 5D에 필요한 8월 28일·31일·9월 1일 종가가 없어, 기존 코드가 8월 21일짜리 옛 완성형 데이터로 자동 후퇴했다.")
    add_tagged(doc, "FACT", "오류는 날짜 표기 문제가 아니라 원천 선택 문제였다. 기존 JSON도 sourceMode=historical_fallback, asOf=2026-08-21, fallbackReason='found 3'을 기록하고 있었다.")
    add_tagged(doc, "FACT", "이번 재계산은 현재 1,635종목 전체의 8/28~9/4 종가를 복원했다. Master와 겹치는 9/2~9/4의 4,905개 가격을 대조한 결과 중앙 절대오차는 0, 최대 오차도 0.0000046% 미만이었다.")
    add_tagged(doc, "INFERENCE", "따라서 사용자의 지적이 맞다. 9월 4일 데이터가 빠진 것이 아니라, 보고서 로직이 부분적으로 존재하는 최신 구간보다 오래된 완성형 구간을 우선한 설계 결함이었다.")
    add_callout(doc, "재발 방지", "off-by-one 세션 판별을 고치고, --as-of cutoff를 추가하며, 더 오래된 fallback은 명시적으로 허용한 경우에만 쓰도록 안전 패치 사본을 만들었다. 운영 원본은 아직 덮어쓰지 않았다.", fill="E2F0D9", accent=GREEN)
    add_heading(doc, "정확한 5D 정의", level=2)
    add_paragraph(doc, "사용한 종가: 2026-08-28, 8/31, 9/1, 9/2, 9/3, 9/4. 5D는 9/4 종가 ÷ 8/28 종가 - 1이다. 주말 carry 값과 9/7 장중 후보는 배제했다.")

    doc.add_page_break()
    add_heading(doc, "TOPIX Capital Goods: 무엇이 실제로 약했나")
    subrows = []
    for label, key in [
        ("방산", "Defense"), ("AI·DC 직접", "AI_DC_Direct"), ("AI 인접", "AI_Adjacent"),
        ("AI 광의", "AI_Combined"), ("기타", "Other"),
    ]:
        item = japan[key]
        subrows.append([
            label, str(item["n"]), pct(float(item["equalWeight"])), pct(float(item["median"])),
            f"{float(item['breadth']) * 100:.1f}%", pct(float(item["capWeight"])),
            f"{float(item['shareOfGrossDownside']) * 100:.1f}%",
        ])
    add_table(doc, ["분석 묶음", "N", "EW", "중앙", "상승", "CW*", "gross↓"], subrows, [3.5, 1.0, 1.8, 1.8, 1.6, 1.8, 2.0], font_size=7.6)
    add_image(doc, charts["subthemes"], 7.0, "그림 3. AI 직접·인접 묶음을 실제 공급망 기능으로 재분해. 분류는 공식 GICS가 아닌 분석용 프록시.")
    add_tagged(doc, "FACT", "광·케이블 -4.91%, 전력·냉각장비 -4.25%, 반도체장비·FA -2.49%가 약했다. 반면 국내 전력망·전기공사 +0.23%, 시공·HVAC -0.44%는 상대적으로 방어적이었다.")
    add_tagged(doc, "INFERENCE", "‘일본 AI Capital Goods 전체 붕괴’보다는 해외 AI 투자와 글로벌 금리에 민감한 장비·케이블·FA가 약하고, 일본 국내 건설·전력망 실행주는 버틴 선택적 조정으로 읽는 편이 정확하다.")

    doc.add_page_break()
    add_heading(doc, "시간축: 9월 2일 금리 충격에 하락이 몰렸다")
    add_image(doc, charts["path"], 7.05, "그림 4. 8/28=0 기준 그룹별 동일가중 누적수익률. 9/2 하루 AI 광의 -3.54%, 방산 -3.42%, 기타 -2.28%.")
    add_table(
        doc,
        ["발표/시장일", "관측", "AI·DC 주가 경로", "판독"],
        [
            ["8/31", "7월 산업생산 +0.1% m/m, 예상 -0.6~-0.7%", "AI 광의 +0.53% 누적", "실물지표는 악재가 아님"],
            ["9/1", "제조업 PMI 54.9; AI·반도체 수요", "AI 광의 +0.43% 누적", "수요 신호는 강함"],
            ["9/2", "Takata 매파 발언; JGB 10Y 3.015%", "하루 -3.54%, 누적 -3.13%", "할인율 충격과 정렬"],
            ["9/3", "채권·주식 반등", "누적 -2.70%", "부분 회복"],
            ["9/4", "금리 안정·기술주 반등", "누적 -2.05%", "장비군은 미회복"],
        ],
        [2.1, 6.2, 4.0, 4.5],
        font_size=7.2,
    )
    add_tagged(doc, "FACT", "9월 2일 TOPIX는 -2.40%, Prime 시장 종목의 92%가 하락했다. JGB 10년물은 3.015%로 1996년 이후 최고, 2년물은 1.86%로 1995년 이후 최고를 기록했다.")
    add_tagged(doc, "INFERENCE", "폭넓은 시장 하락과 같은 날 AI·방산의 동반 급락은 개별 주문 악화보다 공통 할인율·유가·엔화 요인의 설명력을 높인다.")

    doc.add_page_break()
    add_heading(doc, "산업생산과의 인과관계: 일본도 직접 원인으로 보기 어렵다")
    add_heading(doc, "8월 31일 발표된 7월 산업생산", level=2)
    add_tagged(doc, "FACT", "일본 7월 산업생산은 +0.1% m/m로 4개월 연속 증가했고, 시장 예상(-0.6~-0.7%)을 웃돌았다. 전년 대비는 +4.1%였다.")
    add_tagged(doc, "FACT", "생산기계, 화학, 전자부품·디바이스는 증가했고, 반도체 제조장비 수출이 생산기계를 지지했다. 운송장비와 금속제품은 감소했다.")
    add_tagged(doc, "FACT", "제조업체 전망은 8월 +6.4%, 9월 -4.2%였고 METI 평가는 '일진일퇴'였다. 헤드라인은 예상보다 강했지만 회복의 폭이 충분히 넓다고 보기도 어렵다.")
    add_heading(doc, "수요 측 교차확인", level=2)
    add_table(
        doc,
        ["발표일", "지표", "결과", "AI·설비투자 해석"],
        [
            ["8/19", "6월 핵심 기계수주", "+9.7% m/m; 3Q 전망 +4.9%", "설비투자 파이프라인 개선"],
            ["8/27", "BOJ Himino 연설", "미국 IT투자 $1.7tn; hyperscaler $0.8tn", "일본 반도체장비·부품 파급 인정"],
            ["8/31", "7월 산업생산", "+0.1% m/m vs 예상 감소", "반도체장비가 생산기계 지지"],
            ["9/1", "8월 제조업 PMI", "54.9; 신규주문 2018년 이후 최고", "AI·반도체 수요가 명시적 동력"],
            ["9/1", "2Q 법인 설비투자", "+1.6% y/y; ex-SW +2.9% q/q", "예상(-0.3%) 상회"],
        ],
        [2.0, 4.2, 5.0, 5.6],
        font_size=7.3,
    )
    add_tagged(doc, "INFERENCE", "산업생산이 약해서 AI 자본재가 하락했다면, 예상 상회한 산업생산·강한 PMI·기계수주와 설명이 충돌한다. 실물 흐름은 'AI 강세, 비AI 혼재'에 가깝다.")
    add_tagged(doc, "INFERENCE", "역설적으로 AI 수요가 성장·물가 전망을 지지하면 BOJ 정상화 기대와 장기금리를 높여 AI주 멀티플을 압박할 수 있다. 펀더멘털 호재와 단기 주가 악재가 동시에 성립한다.")
    add_tagged(doc, "SPECULATION", "광케이블·전력장비·FA의 9/4 미회복에는 포지셔닝·고밸류 부담이나 해외 기술주 조정이 추가로 작용했을 수 있다. 5일 이벤트 정렬만으로 개별 팩터를 완전히 분리할 수는 없다.")

    doc.add_page_break()
    add_heading(doc, "유럽 재확인: 같은 방향, 다른 하드데이터 시점")
    add_image(doc, EUROPE_MEMO_DIR / "01_group_returns.png", 6.95, "그림 5. STOXX Europe 600 Capital Goods 그룹별 동일·시총가중 5D 수익률.")
    add_tagged(doc, "FACT", "유럽 Capital Goods 100개는 -1.99%, 방산 12개 -6.41%, AI·데이터센터 광의 15개 -2.39%, 기타 73개 -1.18%였다.")
    add_tagged(doc, "FACT", "AI 광의는 STOXX600 전체(-0.9%)보다 1.4%p, Industrials(-1.9%)보다 0.5%p 약했다. 시총가중 하락의 43.1%를 AI 광의가 설명했다.")
    add_tagged(doc, "FACT", "독일 7월 산업생산(-1.1% m/m, 자본재 -3.4%)은 9월 7일 발표돼 9월 4일에 끝난 가격 구간의 직접 촉매가 될 수 없다. 당시 확인 가능했던 유로존 6월 자본재 생산은 -1.4%로 기저는 약했다.")
    add_tagged(doc, "FACT", "반대로 9월 1일 독일 제조업 PMI는 54.3으로 상승했고, 기업들은 방산 지출·데이터센터 건설·재고 축적을 수요 동력으로 언급했다.")
    add_tagged(doc, "INFERENCE", "유럽도 산업생산은 중기 배경이지만 단기 촉매는 아니다. 8/31~9/2 유가·인플레이션·독일 장기금리 상승과 AI 대형주 하락의 시간 정렬이 더 강하다.")
    add_callout(doc, "공통점과 차이", "공통점은 방산의 큰 개별 낙폭과 AI 대형주의 높은 손실 기여다. 차이는 일본의 당일 충격이 9/2 JGB 급등에 더 선명하게 집중됐고, 국내 전력망·시공주는 회복한 반면 케이블·전력/냉각·semicap/FA만 약세를 남겼다는 점이다.", fill="FFF2CC", accent=AMBER)

    doc.add_page_break()
    add_heading(doc, "TOPIX 손실 기여 종목")
    add_image(doc, charts["contributors"], 6.9, "그림 6. 현재 Capital Goods 내 시작 시가총액×수익률의 진단 기여. 공식 TOPIX 기여도와는 다름.")
    top_rows = sorted(
        [row for row in japan_stocks if float(row.get("downsideContributionPp") or 0) < 0],
        key=lambda row: float(row["downsideContributionPp"]),
    )[:10]
    labels = {"Defense": "방산", "AI_DC_Direct": "AI 직접", "AI_Adjacent": "AI 인접", "Other": "기타"}
    add_table(
        doc,
        ["종목", "코드", "분류", "5D", "CG 기여"],
        [[clean_name(row["name"]), row["ticker"].split()[0], labels[row["bucket"]], pct(float(row["return5d"])), f"{float(row['downsideContributionPp']) * 100:.3f}%p"] for row in top_rows],
        [5.6, 1.8, 2.2, 2.2, 2.5],
        font_size=7.6,
    )
    add_tagged(doc, "FACT", "Mitsubishi Heavy가 -0.372%p로 최대 하락기여였고, Mitsubishi Electric -0.344%p, Hitachi -0.319%p, Fujikura -0.207%p가 뒤를 이었다.")
    add_tagged(doc, "INFERENCE", "상위 네 종목 중 세 종목이 AI 직접군이라는 점이 일본 AI 바스켓의 gross downside 61.6%를 설명한다. 평균 낙폭보다 시가총액과 노출의 결합이 중요했다.")

    doc.add_page_break()
    add_heading(doc, "AI 프록시 분류와 대표 사업 근거")
    add_paragraph(doc, "분류는 Bloomberg/GICS 공식 테마가 아니라 Capital Goods 구성종목의 실제 공급망 역할을 기준으로 만든 상호배타적 분석용 프록시다. 방산 중복은 방산을 우선했다.", size=8.8, color=GRAY)
    for title, key, description in [
        ("직접 — 데이터센터 시공·HVAC", "Direct_Construction_HVAC", "데이터센터 MEP·통신·공조 시공 및 운영"),
        ("직접 — 광·케이블", "Direct_Cable", "광섬유·전력케이블 등 AI 데이터센터 연결 인프라"),
        ("직접 — 전력·냉각장비", "Direct_Power_Cooling_Equipment", "UPS·배전·발전·냉각·랙/외함"),
        ("인접 — 전력망·전기공사", "Adjacent_Grid_Electrical", "국내 송배전·전기공사 및 grid 증설"),
        ("인접 — 반도체장비·FA", "Adjacent_Semicap_FA", "웨이퍼 이송·공정 보조·로봇·서보·자동화"),
    ]:
        add_heading(doc, f"{title}", level=2)
        add_paragraph(doc, f"{description}. {names_by_subbucket(japan_stocks, key)}", size=7.6, after=3)
    add_heading(doc, "대표 기업의 공식 설명", level=2)
    source_line(doc, "Fuji Electric — Data center UPS / power systems", "https://www.fujielectric.com/about/example/detail/solution_datacenter.html", "AI 데이터센터용 UPS·전력 안정화")
    source_line(doc, "EBARA — Long-term Vision / data center cooling", "https://ebara.com/content/dam/ebara/grand-masters/entities/en/pdf/ir/business/vision/161_q4_e-plan2028_en.pdf.coredownload.pdf", "데이터센터 냉각용 펌프와 반도체 chiller")
    source_line(doc, "Yaskawa — Industrial robots", "https://www.yaskawa-global.com/product/robotics", "AI 자동화·반도체 제조용 로봇")
    source_line(doc, "DAIHEN — Wafer transfer robots", "https://www.daihen.co.jp/en/products/cleanrobot/wafer/", "반도체 웨이퍼 이송 로봇")
    source_line(doc, "FANUC — Physical AI collaboration", "https://www.fanuc.co.jp/en/profile/pr/newsrelease/2026/notice20260513.html", "Physical AI·로봇 자동화")

    doc.add_page_break()
    add_heading(doc, "Risks, 확인할 지표, 최종 판단")
    add_heading(doc, "결론을 바꿀 수 있는 반증", level=2)
    add_paragraph(doc, "금리가 안정된 뒤에도 Fujikura·Mitsubishi Electric·Hitachi·Ebara·Yaskawa·SMC가 계속 시장을 하회하고, 데이터센터·반도체장비 주문과 백로그가 둔화한다면 수요 우려의 비중을 높여야 한다.", bullet=True)
    add_paragraph(doc, "일본 8월 산업생산 전망(+6.4%)이 실제치에서 크게 하향되고 생산기계·전자부품이 꺾이면 현재의 '할인율 1차, 실물 2차' 판단은 약해진다.", bullet=True)
    add_paragraph(doc, "BOJ의 9월 정책 경로, JGB 2년·10년 금리, 엔화, 유가가 안정되는데도 AI 장비주가 회복하지 못하면 포지셔닝보다 이익추정치 리스크를 우선해야 한다.", bullet=True)
    add_heading(doc, "데이터 한계", level=2)
    add_tagged(doc, "FACT", "현지가격 수익률이라 배당과 FX는 제외된다. 현재 구성종목을 과거에 소급하므로 생존편향이 있다. 일본 시총은 9/2 mcap과 주가비율로 8/28을 역산했다.")
    add_tagged(doc, "INFERENCE", "5일 이벤트 정렬은 인과의 필요조건을 점검하지만 정식 팩터 회귀는 아니다. 본 보고서의 인과 판정은 '산업생산 직접 원인 낮음, 할인율·금리 충격 높음' 수준으로 읽어야 한다.")
    add_callout(doc, "최종 판단", "유럽과 일본 모두 AI 관련 Capital Goods는 기타보다 나빴다. 일본은 AI 광의가 Capital Goods gross downside의 61.6%를 차지했다. 그러나 7월 산업생산은 일본에서 예상 상회했고 유럽에서는 가격 구간 뒤 발표됐다. 공통된 단기 촉매는 장기금리·유가·인플레이션이며, 산업생산은 비AI 제조업의 혼재된 중기 배경으로 보는 것이 가장 일관적이다.")

    doc.add_page_break()
    add_heading(doc, "Sources & Method")
    add_heading(doc, "시장 데이터와 감사 추적", level=2)
    add_paragraph(doc, f"유럽: {EUROPE_JSON} · STOXX600 SXXP Raw · 2026-08-28~09-04", size=7.5)
    add_paragraph(doc, f"일본: {JAPAN_JSON} · 현재 TOPIX 1,635종목 중 Capital Goods 293개 · 2026-08-28~09-04", size=7.5)
    add_paragraph(doc, "TOPIX 공개종가 1,635/1,635 확보, 로컬 BQL 9/2~9/4 겹침 4,905건 대조. 중앙 절대오차 0, 최대 4.60e-08(비율).", size=7.5)
    add_heading(doc, "일본 거시·시장", level=2)
    source_line(doc, "METI — Indices of Industrial Production", "https://www.meti.go.jp/english/statistics/tyo/iip/", "7월 산업생산 공식 통계")
    source_line(doc, "Cabinet Office — June machinery orders", "https://www.esri.cao.go.jp/en/stat/juchu/2026/2606juchu-e.html", "핵심 수주 +9.7% m/m, 3Q 전망 +4.9%")
    source_line(doc, "BOJ Himino speech (2026-08-27)", "https://www.boj.or.jp/en/about/press/koen_2026/ko260827a.htm", "AI 수요의 성장·물가 파급")
    source_line(doc, "BOJ Takata speech (2026-09-02)", "https://www.boj.or.jp/en/about/press/koen_2026/ko260902a.htm", "통화정책 정상화 발언")
    source_line(doc, "Reuters — BOJ must conduct rate hikes nimbly", "https://www.investing.com/news/economy-news/boj-must-conduct-rate-hikes-nimbly-hawkish-board-member-says-4885257", "9/2 발언과 정책 배경")
    source_line(doc, "Reuters — Japan stocks, oil and bond yields", "https://www.indopremier.com/module/newsDetail.php?group_news=IPOTNEWS&halaman=1&jdl=Japan_stocks_end_lower_as_US_Iran_strikes_lift_oil__bond_yields&name=&news_id=237677&q=japan+shares,+nikkei+225,+topix,&search=y_general&taging_subtype=MARKETOVERVIEW", "TOPIX -2.4%, JGB 10Y 3.015%")
    source_line(doc, "Reuters — Japan manufacturing PMI", "https://currently.att.yahoo.com/att/japan-manufacturing-business-rises-fastest-003441616.html", "PMI 54.9, AI·반도체 수요")
    source_line(doc, "TOPIX historical prices", "https://ca.investing.com/indices/topix-historical-data", "8/28 4,146.71 → 9/4 4,103.23")
    add_heading(doc, "유럽 거시·시장", level=2)
    source_line(doc, "Destatis — German July industrial production", "https://www.destatis.de/EN/Press/2026/09/PE26_316_421.html", "9/7 발표, 전체 -1.1%, 자본재 -3.4%")
    source_line(doc, "Eurostat — Euro area June industrial production", "https://ec.europa.eu/eurostat/web/products-euro-indicators/w/4-13082026-ap", "전체 0.0%, 자본재 -1.4%")
    source_line(doc, "Reuters — German August manufacturing PMI", "https://www.marketscreener.com/news/german-manufacturing-growth-accelerates-in-august-as-orders-surge-pmi-shows-ce7858dddc80fe2d", "PMI 54.3, 데이터센터 건설 언급")
    source_line(doc, "ifo — German export expectations", "https://www.ifo.de/en/press-release/2026-08-26/export-expectations-germany-rise-sharply-august-2026", "전기장비·전자광학 낙관, AI 투자")
    source_line(doc, "Reuters — European yields and inflation", "https://www.marketscreener.com/news/european-shares-hit-by-rising-bond-yields-energy-driven-inflation-concerns-ce7858d3db89f72c", "독일 10년물 2011년 이후 최고")

    core = doc.core_properties
    core.title = "유럽·일본 Capital Goods 5D 비교: AI 약세와 산업생산 인과관계"
    core.subject = "STOXX Europe 600 and TOPIX capital goods attribution"
    core.author = "OpenAI Codex"
    core.comments = "Data as of 2026-09-04; prepared 2026-09-08"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


def main() -> None:
    europe, japan, japan_meta, japan_stocks = load_data()
    charts = make_charts(europe, japan, japan_stocks)
    build_document(europe, japan, japan_meta, japan_stocks, charts)


if __name__ == "__main__":
    main()
