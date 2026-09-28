from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl
from PIL import Image, ImageDraw, ImageFont


HERE = Path(__file__).resolve().parent
BQL_DIR = HERE.parents[3]
DEFAULT_MASTER = BQL_DIR / "Rawfile" / "BQuant_Master.xlsx"
DEFAULT_OUTPUT_ROOT = BQL_DIR / "Theme" / "output"
DEFAULT_MATERIAL_ROOT = BQL_DIR / "Theme" / "Material" / "Sector_Industry"
DEFAULT_TOPIX_FALLBACK = (
    Path.home()
    / "Documents"
    / "python"
    / "모닝미팅_리서치발표"
    / "새 폴더"
    / "BQL_Market_Rotation_Transfer_20260901"
    / "raw"
    / "consolidated_reference"
    / "Bquant_All_19_Grouped_Raw.xlsx"
)

MARKETS = {
    "STOXX600": {
        "key": "STOXX600",
        "title": "STOXX Europe 600",
        "sheet": "SXXP Raw",
        "source_code": "SXXP Index",
        "currency": "Mixed local currencies",
        "chart_prefix": "stoxx600",
    },
    "TOPIX": {
        "key": "TOPIX",
        "title": "TOPIX",
        "sheet": "TPX500 Raw",
        "source_code": "TPX500 Index",
        "currency": "JPY",
        "chart_prefix": "topix",
    },
}

METRICS = [
    "12MF P/E Blended",
    "12MF EPS Blended",
    "FY1 EPS",
    "FY2 EPS",
    "Price",
    "Market Cap (Daily)",
    "Current ROE",
    "Current Operating Margin",
    "12MF ROE",
    "12MF Operating Margin",
]

COLORS = {
    "navy": "#17365D",
    "blue": "#2F75B5",
    "pale": "#EAF2F8",
    "red": "#C00000",
    "green": "#008000",
    "gray": "#666666",
    "light_gray": "#D9E2F3",
    "grid": "#D9D9D9",
    "black": "#111111",
    "white": "#FFFFFF",
}


@dataclass
class MarketPanel:
    key: str
    title: str
    source_sheet: str
    source_code: str
    currency: str
    source_workbook: str
    source_mode: str
    tickers: list[str]
    names: list[str]
    sectors: list[str]
    industries: list[str]
    dates: list[date]
    prices: np.ndarray
    market_caps: np.ndarray


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build five-session STOXX Europe 600 and TOPIX sector/industry datasets and charts."
    )
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    parser.add_argument(
        "--output-root",
        "--output-dir",
        dest="output_root",
        type=Path,
        default=DEFAULT_OUTPUT_ROOT,
        help="Date-first result root. Charts are saved to <root>/<YYYYMMDD>/Sector_Industry.",
    )
    parser.add_argument(
        "--material-root",
        type=Path,
        default=DEFAULT_MATERIAL_ROOT,
        help="Calculation-data root. JSON is saved to <root>/<YYYYMMDD>/data.",
    )
    parser.add_argument(
        "--topix-fallback",
        type=Path,
        default=DEFAULT_TOPIX_FALLBACK,
        help="Long-format historical workbook used only when the Master lacks six valid TOPIX closes.",
    )
    parser.add_argument(
        "--market",
        choices=["ALL", "STOXX600", "TOPIX"],
        default="ALL",
        help="Market to build. ALL creates both markets in one JSON file.",
    )
    parser.add_argument("--top-n", type=int, default=8)
    return parser.parse_args()


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)) and 20_000 < float(value) < 80_000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    return None


def number(value: Any) -> float:
    if isinstance(value, bool):
        return np.nan
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return np.nan


def clean_group(value: Any, fallback: str = "Unclassified") -> str:
    text = str(value or "").strip()
    return text if text else fallback


def load_metric_block(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    anchor_row: int,
    ticker_to_row: dict[str, int],
    member_count: int,
) -> tuple[list[date], np.ndarray]:
    label = ws.cell(anchor_row, 7).value
    if not isinstance(label, str) or " | " not in label:
        raise ValueError(f"Metric anchor not found at {ws.title}!G{anchor_row}: {label!r}")

    header = next(
        ws.iter_rows(
            min_row=anchor_row + 1,
            max_row=anchor_row + 1,
            min_col=7,
            max_col=ws.max_column,
            values_only=True,
        )
    )
    date_columns: list[tuple[int, date]] = []
    for offset, value in enumerate(header):
        parsed = as_date(value)
        if parsed is not None:
            date_columns.append((offset, parsed))
    if not date_columns:
        raise ValueError(f"No observation dates in {ws.title} metric block {label}")

    dates = [value for _, value in date_columns]
    matrix = np.full((member_count, len(dates)), np.nan, dtype=np.float64)
    rows = ws.iter_rows(
        min_row=anchor_row + 2,
        max_row=anchor_row + 1 + member_count,
        min_col=7,
        max_col=ws.max_column,
        values_only=True,
    )
    for row in rows:
        ticker = row[0]
        if not isinstance(ticker, str):
            continue
        out_row = ticker_to_row.get(ticker.strip())
        if out_row is None:
            continue
        for out_col, (offset, _) in enumerate(date_columns):
            if offset < len(row):
                matrix[out_row, out_col] = number(row[offset])
    return dates, matrix


def align_matrix(source_dates: list[date], source: np.ndarray, target_dates: list[date]) -> np.ndarray:
    if source_dates == target_dates:
        return source
    lookup = {value: idx for idx, value in enumerate(source_dates)}
    result = np.full((source.shape[0], len(target_dates)), np.nan, dtype=np.float64)
    for out_col, value in enumerate(target_dates):
        source_col = lookup.get(value)
        if source_col is not None:
            result[:, out_col] = source[:, source_col]
    return result


def load_market_panel(master: Path, config: dict[str, str]) -> MarketPanel:
    workbook = openpyxl.load_workbook(master, read_only=True, data_only=True)
    try:
        if config["sheet"] not in workbook.sheetnames:
            raise KeyError(f"Missing sheet: {config['sheet']}")
        ws = workbook[config["sheet"]]
        live_value = ws.cell(3, 6).value
        if not isinstance(live_value, (int, float)) or int(live_value) <= 0:
            raise ValueError(f"Invalid live member count in {config['sheet']}!F3: {live_value!r}")
        member_count = int(live_value)

        tickers: list[str] = []
        names: list[str] = []
        sectors: list[str] = []
        industries: list[str] = []
        for row in ws.iter_rows(
            min_row=8,
            max_row=7 + member_count,
            min_col=1,
            max_col=4,
            values_only=True,
        ):
            ticker = str(row[0] or "").strip()
            if not ticker:
                continue
            tickers.append(ticker)
            names.append(str(row[1] or ticker).strip())
            sectors.append(clean_group(row[2]))
            industries.append(clean_group(row[3]))

        if len(tickers) != member_count:
            raise ValueError(
                f"Expected {member_count} metadata rows in {config['sheet']}, found {len(tickers)}"
            )
        if len(set(tickers)) != len(tickers):
            raise ValueError(f"Duplicate tickers in {config['sheet']} current membership block")

        ticker_to_row = {ticker: idx for idx, ticker in enumerate(tickers)}
        first_anchor = member_count + 10
        block_stride = member_count + 2
        price_anchor = first_anchor + METRICS.index("Price") * block_stride
        mcap_anchor = first_anchor + METRICS.index("Market Cap (Daily)") * block_stride

        price_dates, prices = load_metric_block(ws, price_anchor, ticker_to_row, member_count)
        mcap_dates, market_caps = load_metric_block(ws, mcap_anchor, ticker_to_row, member_count)
        market_caps = align_matrix(mcap_dates, market_caps, price_dates)

        source_code = str(ws.cell(4, 2).value or config["source_code"]).strip()
        currency = str(ws.cell(4, 4).value or config["currency"]).strip()
        return MarketPanel(
            key=config["key"],
            title=config["title"],
            source_sheet=config["sheet"],
            source_code=source_code,
            currency=currency,
            source_workbook=str(master.resolve()),
            source_mode="master",
            tickers=tickers,
            names=names,
            sectors=sectors,
            industries=industries,
            dates=price_dates,
            prices=prices,
            market_caps=market_caps,
        )
    finally:
        workbook.close()


def load_topix_fallback_panel(source: Path) -> MarketPanel:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if "Stock_Daily" not in workbook.sheetnames:
            raise KeyError(f"Missing Stock_Daily sheet in {source}")
        ws = workbook["Stock_Daily"]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = {str(value): idx for idx, value in enumerate(header) if value is not None}
        required = [
            "Date",
            "Index_Code",
            "Ticker",
            "Security_Name",
            "GICS_Sector",
            "GICS_Industry_Group",
            "Price",
            "MarketCap",
        ]
        missing = [name for name in required if name not in columns]
        if missing:
            raise KeyError(f"Missing columns in Stock_Daily: {missing}")

        observations: dict[date, dict[str, tuple[str, str, str, float, float]]] = defaultdict(dict)
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = str(row[columns["Index_Code"]] or "").strip().upper()
            if code not in {"TPX", "TPX500"}:
                continue
            obs_date = as_date(row[columns["Date"]])
            ticker = str(row[columns["Ticker"]] or "").strip()
            if obs_date is None or not ticker:
                continue
            observations[obs_date][ticker] = (
                str(row[columns["Security_Name"]] or ticker).strip(),
                clean_group(row[columns["GICS_Sector"]]),
                clean_group(row[columns["GICS_Industry_Group"]]),
                number(row[columns["Price"]]),
                number(row[columns["MarketCap"]]),
            )
        if not observations:
            raise ValueError(f"No TOPIX observations in {source}")

        dates = sorted(observations)
        latest_date = dates[-1]
        latest = observations[latest_date]
        tickers = sorted(latest)
        ticker_to_row = {ticker: idx for idx, ticker in enumerate(tickers)}
        names = [latest[ticker][0] for ticker in tickers]
        sectors = [latest[ticker][1] for ticker in tickers]
        industries = [latest[ticker][2] for ticker in tickers]
        prices = np.full((len(tickers), len(dates)), np.nan, dtype=np.float64)
        market_caps = np.full_like(prices, np.nan)
        for out_col, obs_date in enumerate(dates):
            for ticker, values in observations[obs_date].items():
                out_row = ticker_to_row.get(ticker)
                if out_row is None:
                    continue
                prices[out_row, out_col] = values[3]
                market_caps[out_row, out_col] = values[4]

        return MarketPanel(
            key="TOPIX",
            title="TOPIX",
            source_sheet="Stock_Daily",
            source_code="TPX historical fallback",
            currency="JPY",
            source_workbook=str(source.resolve()),
            source_mode="historical_fallback",
            tickers=tickers,
            names=names,
            sectors=sectors,
            industries=industries,
            dates=dates,
            prices=prices,
            market_caps=market_caps,
        )
    finally:
        workbook.close()


def actual_sessions(prices: np.ndarray) -> list[int]:
    sessions: list[int] = []
    for current in range(1, prices.shape[1]):
        before = prices[:, current - 1]
        after = prices[:, current]
        valid = np.isfinite(before) & np.isfinite(after) & (before > 0) & (after > 0)
        if int(valid.sum()) < 10:
            continue
        changed = ~np.isclose(after[valid], before[valid], rtol=1e-7, atol=1e-9)
        if float(changed.mean()) > 0.10:
            sessions.append(current)
    return sessions


def safe_ratio(after: np.ndarray, before: np.ndarray) -> np.ndarray:
    result = np.full(before.shape, np.nan, dtype=np.float64)
    valid = np.isfinite(before) & np.isfinite(after) & (before > 0) & (after > 0)
    result[valid] = after[valid] / before[valid] - 1.0
    return result


def json_number(value: float | None) -> float | None:
    if value is None or not math.isfinite(float(value)):
        return None
    return float(value)


def percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    return float(np.percentile(np.asarray(values, dtype=float), q))


def summarize(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(records)
    values = [float(row["return5d"]) for row in rows if row.get("return5d") is not None and not row["excluded"]]
    positive = sum(value > 0 for value in values)
    return {
        "members": len(rows),
        "validMembers": len(values),
        "coverage": json_number(len(values) / len(rows)) if rows else None,
        "return5d": json_number(float(np.mean(values))) if values else None,
        "median5d": json_number(float(np.median(values))) if values else None,
        "breadth5d": json_number(positive / len(values)) if values else None,
        "dispersion5d": json_number(float(np.std(values, ddof=0))) if values else None,
        "p25": json_number(percentile(values, 25)),
        "p75": json_number(percentile(values, 75)),
    }


def stock_slice(records: list[dict[str, Any]], top_n: int) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    valid = [row for row in records if row.get("return5d") is not None and not row["excluded"]]
    ranked = sorted(valid, key=lambda row: float(row["return5d"]), reverse=True)
    gainers = [row for row in ranked if float(row["return5d"]) > 0]
    losers = [row for row in reversed(ranked) if float(row["return5d"]) < 0]
    top = gainers[: min(top_n, len(gainers))]
    bottom = losers[: min(top_n, len(losers))]
    fields = ("ticker", "name", "industry", "return5d", "marketCapReturn5d")
    compact = lambda row: {key: row.get(key) for key in fields}
    return [compact(row) for row in top], [compact(row) for row in bottom]


def leadership_state(sector_return: float | None, benchmark: float | None, breadth: float | None) -> str:
    if sector_return is None or benchmark is None or breadth is None:
        return "판정 유보"
    above = sector_return >= benchmark
    broad = breadth >= 0.50
    if above and broad:
        return "강세 확산"
    if above and not broad:
        return "상대 강세 집중"
    if not above and broad:
        return "방어적 확산"
    return "약세 확산"


def corporate_action_reason(price_return: float, mcap_return: float | None) -> str | None:
    if price_return <= -0.90 or price_return >= 3.00:
        return "extreme_price_move"
    if mcap_return is None or 1.0 + price_return <= 0 or 1.0 + mcap_return <= 0:
        return None
    log_gap = abs(math.log1p(price_return) - math.log1p(mcap_return))
    if abs(price_return) >= 0.35 and log_gap >= 0.25:
        return "possible_corporate_action"
    return None


def build_market(panel: MarketPanel, top_n: int) -> dict[str, Any]:
    sessions = actual_sessions(panel.prices)
    if len(sessions) < 6:
        raise ValueError(f"{panel.title}: at least six valid closes are required; found {len(sessions)}")
    end_idx = sessions[-1]
    start_idx = sessions[-6]

    price_returns = safe_ratio(panel.prices[:, end_idx], panel.prices[:, start_idx])
    mcap_returns = safe_ratio(panel.market_caps[:, end_idx], panel.market_caps[:, start_idx])

    stocks: list[dict[str, Any]] = []
    for idx, ticker in enumerate(panel.tickers):
        price_return = json_number(price_returns[idx])
        mcap_return = json_number(mcap_returns[idx])
        reason = corporate_action_reason(price_return, mcap_return) if price_return is not None else "missing_price"
        stocks.append(
            {
                "ticker": ticker,
                "name": panel.names[idx],
                "sector": panel.sectors[idx],
                "industry": panel.industries[idx],
                "startPrice": json_number(panel.prices[idx, start_idx]),
                "endPrice": json_number(panel.prices[idx, end_idx]),
                "return5d": price_return,
                "startMarketCap": json_number(panel.market_caps[idx, start_idx]),
                "endMarketCap": json_number(panel.market_caps[idx, end_idx]),
                "marketCapReturn5d": mcap_return,
                "excluded": reason is not None,
                "exclusionReason": reason,
            }
        )

    overall = summarize(stocks)
    benchmark = overall["return5d"]
    sector_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    industry_rows: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in stocks:
        sector_rows[row["sector"]].append(row)
        industry_rows[(row["sector"], row["industry"])].append(row)

    industries: list[dict[str, Any]] = []
    for (sector, industry), rows in industry_rows.items():
        summary = summarize(rows)
        summary.update({"sector": sector, "industry": industry})
        summary["excessVsMarket"] = json_number(summary["return5d"] - benchmark) if summary["return5d"] is not None and benchmark is not None else None
        industries.append(summary)

    sector_summaries: list[dict[str, Any]] = []
    for sector, rows in sector_rows.items():
        summary = summarize(rows)
        sector_industries = [row for row in industries if row["sector"] == sector and row["return5d"] is not None]
        sector_industries.sort(key=lambda row: float(row["return5d"]), reverse=True)
        top_gainers, bottom_losers = stock_slice(rows, top_n)
        summary.update(
            {
                "sector": sector,
                "industryCount": len(sector_industries),
                "excessVsMarket": json_number(summary["return5d"] - benchmark) if summary["return5d"] is not None and benchmark is not None else None,
                "state": leadership_state(summary["return5d"], benchmark, summary["breadth5d"]),
                "topIndustry": sector_industries[0] if sector_industries else None,
                "bottomIndustry": sector_industries[-1] if sector_industries else None,
                "industrySpread": json_number(sector_industries[0]["return5d"] - sector_industries[-1]["return5d"])
                if len(sector_industries) >= 2
                else 0.0 if len(sector_industries) == 1 else None,
                "topGainers": top_gainers,
                "bottomLosers": bottom_losers,
            }
        )
        sector_summaries.append(summary)

    sector_summaries.sort(
        key=lambda row: float(row["return5d"]) if row["return5d"] is not None else -math.inf,
        reverse=True,
    )
    industries.sort(
        key=lambda row: float(row["return5d"]) if row["return5d"] is not None else -math.inf,
        reverse=True,
    )
    for rank, row in enumerate(sector_summaries, 1):
        row["rank"] = rank
    for sector, rows in sector_rows.items():
        valid = [row for row in rows if row.get("return5d") is not None and not row["excluded"]]
        valid.sort(key=lambda row: float(row["return5d"]), reverse=True)
        for rank, row in enumerate(valid, 1):
            row["sectorRank"] = rank

    return {
        "key": panel.key,
        "title": panel.title,
        "sourceSheet": panel.source_sheet,
        "sourceIndexLabel": panel.source_code,
        "currency": panel.currency,
        "sourceWorkbook": panel.source_workbook,
        "sourceMode": panel.source_mode,
        "asOf": panel.dates[end_idx].isoformat(),
        "start5d": panel.dates[start_idx].isoformat(),
        "sessionDates": [panel.dates[idx].isoformat() for idx in sessions[-6:]],
        "memberCount": len(stocks),
        "validMemberCount": overall["validMembers"],
        "excludedCount": sum(row["excluded"] for row in stocks),
        "overall": overall,
        "sectors": sector_summaries,
        "industries": industries,
        "stocks": stocks,
        "charts": {},
    }


def font_path(bold: bool = False) -> Path | None:
    candidates = [
        Path("C:/Windows/Fonts/malgunbd.ttf" if bold else "C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    return next((path for path in candidates if path.exists()), None)


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = font_path(bold)
    return ImageFont.truetype(str(path), size=size) if path else ImageFont.load_default()


def text_width(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> int:
    box = draw.textbbox((0, 0), text, font=font)
    return int(box[2] - box[0])


def fit_label(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return text[: max(1, limit - 1)].rstrip() + "…"


def draw_horizontal_bar_chart(
    items: list[dict[str, Any]],
    label_key: str,
    value_key: str,
    title: str,
    subtitle: str,
    output: Path,
    benchmark: float | None = None,
    percent_axis: bool = True,
    sort_desc: bool = True,
) -> None:
    clean = [row for row in items if row.get(value_key) is not None]
    clean.sort(key=lambda row: float(row[value_key]), reverse=sort_desc)
    if not clean:
        return

    row_height = 68
    width = 1800
    height = max(700, 245 + row_height * len(clean))
    image = Image.new("RGB", (width, height), COLORS["white"])
    draw = ImageDraw.Draw(image)
    title_font = get_font(48, bold=True)
    subtitle_font = get_font(25)
    label_font = get_font(27)
    value_font = get_font(25, bold=True)
    tick_font = get_font(20)

    draw.text((70, 45), title, fill=COLORS["black"], font=title_font)
    draw.text((70, 112), subtitle, fill=COLORS["gray"], font=subtitle_font)

    chart_left, chart_right = 570, 1695
    chart_top = 205
    values = [float(row[value_key]) for row in clean]
    if percent_axis:
        lo = min(values + ([benchmark] if benchmark is not None else []) + [0.0])
        hi = max(values + ([benchmark] if benchmark is not None else []) + [0.0])
        span = max(hi - lo, 0.01)
        lo -= span * 0.14
        hi += span * 0.14
    else:
        lo, hi = 0.0, 1.0

    def x_for(value: float) -> int:
        return int(chart_left + (value - lo) / (hi - lo) * (chart_right - chart_left))

    for frac in np.linspace(0, 1, 5):
        value = lo + (hi - lo) * float(frac)
        x = x_for(value)
        draw.line((x, chart_top - 18, x, height - 70), fill=COLORS["grid"], width=2)
        label = f"{value * 100:.1f}%" if percent_axis else f"{value * 100:.0f}%"
        draw.text((x - text_width(draw, label, tick_font) // 2, height - 55), label, fill=COLORS["gray"], font=tick_font)

    zero_x = x_for(0.0)
    draw.line((zero_x, chart_top - 20, zero_x, height - 70), fill=COLORS["gray"], width=4)
    if benchmark is not None:
        bench_x = x_for(float(benchmark))
        for y in range(chart_top - 15, height - 70, 18):
            draw.line((bench_x, y, bench_x, min(y + 10, height - 70)), fill=COLORS["navy"], width=4)

    for idx, row in enumerate(clean):
        value = float(row[value_key])
        y = chart_top + idx * row_height
        label = fit_label(str(row[label_key]), 31)
        draw.text((70, y + 10), label, fill=COLORS["black"], font=label_font)
        value_x = x_for(value)
        left, right = sorted((zero_x, value_x))
        if right - left < 3:
            right = left + 3
        color = COLORS["blue"] if value >= 0 else COLORS["red"]
        draw.rounded_rectangle((left, y + 6, right, y + 48), radius=8, fill=color)
        value_label = f"{value * 100:+.1f}%" if percent_axis else f"{value * 100:.0f}%"
        if value >= 0:
            tx = min(value_x + 12, chart_right - text_width(draw, value_label, value_font))
        else:
            tx = max(chart_left, value_x - 12 - text_width(draw, value_label, value_font))
        draw.text((tx, y + 10), value_label, fill=color, font=value_font)

    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, format="PNG", optimize=True)


def slug(value: str) -> str:
    ascii_only = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_").lower()
    return ascii_only or "group"


def create_charts(market: dict[str, Any], chart_dir: Path) -> None:
    prefix = MARKETS[market["key"]]["chart_prefix"]
    tag = market["asOf"].replace("-", "")
    sector_return = chart_dir / f"{prefix}_sector_return_5d_{tag}.png"
    sector_breadth = chart_dir / f"{prefix}_sector_breadth_5d_{tag}.png"
    industry_rotation = chart_dir / f"{prefix}_industry_rotation_5d_{tag}.png"

    draw_horizontal_bar_chart(
        market["sectors"],
        "sector",
        "return5d",
        f"{market['title']} 섹터별 최근 5거래일 수익률",
        f"{market['start5d']} 종가 대비 {market['asOf']} 종가 · 단순수익률 동일가중 평균 · 점선은 시장 전체",
        sector_return,
        benchmark=market["overall"]["return5d"],
    )
    draw_horizontal_bar_chart(
        market["sectors"],
        "sector",
        "breadth5d",
        f"{market['title']} 섹터별 상승 종목 비율",
        "유효 종목 중 최근 5거래일 수익률이 플러스인 비율",
        sector_breadth,
        benchmark=market["overall"]["breadth5d"],
        percent_axis=False,
    )

    industry_valid = [row for row in market["industries"] if row.get("return5d") is not None]
    leaders = industry_valid[:8]
    laggards = list(reversed(industry_valid[-8:]))
    selected = leaders + [row for row in laggards if row not in leaders]
    draw_horizontal_bar_chart(
        selected,
        "industry",
        "return5d",
        f"{market['title']} 산업그룹 상위와 하위",
        "GICS Industry Group 최근 5거래일 동일가중 평균 · 상위 8개와 하위 8개",
        industry_rotation,
        benchmark=market["overall"]["return5d"],
    )

    market["charts"] = {
        "sectorReturn": str(sector_return.resolve()),
        "sectorBreadth": str(sector_breadth.resolve()),
        "industryRotation": str(industry_rotation.resolve()),
        "sectorIndustries": {},
    }
    for sector in market["sectors"]:
        rows = [row for row in market["industries"] if row["sector"] == sector["sector"]]
        path = chart_dir / f"{prefix}_{slug(sector['sector'])}_industries_5d_{tag}.png"
        draw_horizontal_bar_chart(
            rows,
            "industry",
            "return5d",
            f"{sector['sector']} 산업그룹 성과",
            f"{market['title']} · {market['start5d']}~{market['asOf']} · 섹터 평균 {sector['return5d'] * 100:+.1f}%",
            path,
            benchmark=sector["return5d"],
        )
        market["charts"]["sectorIndustries"][sector["sector"]] = str(path.resolve())


def validate_market(market: dict[str, Any], top_n: int) -> None:
    if market["memberCount"] != len(market["stocks"]):
        raise AssertionError(f"{market['key']}: member count mismatch")
    if len({row["ticker"] for row in market["stocks"]}) != market["memberCount"]:
        raise AssertionError(f"{market['key']}: duplicate tickers")
    sector_total = sum(row["members"] for row in market["sectors"])
    industry_total = sum(row["members"] for row in market["industries"])
    if sector_total != market["memberCount"] or industry_total != market["memberCount"]:
        raise AssertionError(f"{market['key']}: hierarchy totals do not reconcile")
    for sector in market["sectors"]:
        if len(sector["topGainers"]) > top_n or len(sector["bottomLosers"]) > top_n:
            raise AssertionError(f"{market['key']} {sector['sector']}: top/bottom limit exceeded")
        overlap = {row["ticker"] for row in sector["topGainers"]} & {row["ticker"] for row in sector["bottomLosers"]}
        if overlap:
            raise AssertionError(f"{market['key']} {sector['sector']}: top/bottom overlap {overlap}")
    if len(market["sessionDates"]) != 6:
        raise AssertionError(f"{market['key']}: five-session window must contain six closes")


def main() -> None:
    args = parse_args()
    master = args.master.resolve()
    output_root = args.output_root.resolve()
    material_root = args.material_root.resolve()
    if not master.exists():
        raise FileNotFoundError(master)
    if not 1 <= args.top_n <= 8:
        raise ValueError("--top-n must be between 1 and 8")

    selected = list(MARKETS) if args.market == "ALL" else [args.market]
    markets: list[dict[str, Any]] = []
    for key in selected:
        panel = load_market_panel(master, MARKETS[key])
        try:
            market = build_market(panel, args.top_n)
        except ValueError as error:
            if key != "TOPIX" or "at least six valid closes" not in str(error):
                raise
            fallback = args.topix_fallback.resolve()
            if not fallback.exists():
                raise ValueError(
                    f"{error}. TOPIX fallback workbook also does not exist: {fallback}"
                ) from error
            panel = load_topix_fallback_panel(fallback)
            market = build_market(panel, args.top_n)
            market["fallbackReason"] = str(error)
        validate_market(market, args.top_n)
        markets.append(market)

    as_of_tag = max(market["asOf"] for market in markets).replace("-", "")
    output_dir = output_root / as_of_tag / "Sector_Industry"
    material_dir = material_root / as_of_tag
    for category in ("AI", "Sector_Industry", "Top10_Down10"):
        (output_root / as_of_tag / category).mkdir(parents=True, exist_ok=True)
    for market in markets:
        chart_dir = output_dir / "charts" / market["key"]
        create_charts(market, chart_dir)
    payload = {
        "meta": {
            "schemaVersion": 1,
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "sourceWorkbook": str(master),
            "topixFallbackWorkbook": str(args.topix_fallback.resolve()),
            "period": "5D",
            "periodDefinition": "latest valid close divided by the close five valid trading sessions earlier minus one",
            "aggregation": "equal-weight arithmetic mean of valid constituent local-price returns",
            "hierarchy": "GICS Sector > GICS Industry Group > constituent",
            "topBottomLimit": args.top_n,
            "notes": [
                "Returns exclude dividends and FX translation.",
                "Current constituents are backfilled through the observation window, so survivorship bias remains.",
                "Obvious corporate-action-like price jumps are excluded when price and market-cap continuity diverge.",
                "The source workbook labels the Japanese sheet TPX500 Raw while the live member count represents the broader TOPIX universe.",
            ],
        },
        "markets": markets,
    }

    data_dir = material_dir / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    output_path = data_dir / f"STOXX600_TOPIX_Sector_Industry_5D_{as_of_tag}.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output_path)


if __name__ == "__main__":
    main()
