from __future__ import annotations

import argparse
import json
import math
import os
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl
from PIL import Image, ImageDraw, ImageFont


HERE = Path(__file__).resolve().parent
# This default is correct after this file is installed in
# BQL\Theme\Material\Sector_Industry\Reports.
BQL_DIR = HERE.parents[3]
DEFAULT_MASTER = BQL_DIR / "Rawfile" / "BQuant_Master.xlsx"
DEFAULT_OUTPUT_ROOT = BQL_DIR / "Theme" / "output"
DEFAULT_MATERIAL_ROOT = BQL_DIR / "Theme" / "Material" / "Sector_Industry"

MARKETS = {
    "STOXX600": {
        "key": "STOXX600",
        "title": "STOXX Europe 600",
        "sheet": "SXXP Raw",
        "source_code": "SXXP Index",
        "currency": "Mixed local currencies",
    },
    "TOPIX": {
        "key": "TOPIX",
        "title": "TOPIX",
        "sheet": "TPX500 Raw",
        "source_code": "TPX500 Index",
        "currency": "JPY",
    },
}

METRICS = [
    "12MF P/E Blended",
    "12MF EPS Blended",
    "FY1 EPS",
    "FY2 EPS",
    "Price",
    "Market Cap (Daily)",
    "Current ROE",
    "Current Operating Margin",
    "12MF ROE",
    "12MF Operating Margin",
]

COLORS = {
    "navy": "#17365D",
    "alternate": "#F3F6F9",
    "grid": "#C8CDD3",
    "red": "#C00000",
    "green": "#008000",
    "black": "#111111",
    "white": "#FFFFFF",
}


@dataclass
class MarketPanel:
    key: str
    title: str
    source_sheet: str
    source_code: str
    currency: str
    source_workbook: str
    tickers: list[str]
    names: list[str]
    sectors: list[str]
    industries: list[str]
    dates: list[date]
    prices: np.ndarray
    market_caps: np.ndarray


def parse_iso_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise argparse.ArgumentTypeError(f"Expected YYYY-MM-DD, received {value!r}") from error


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build 1D or 5D STOXX Europe 600 and TOPIX sector table images and JSON."
    )
    parser.add_argument("--period", choices=["1D", "5D"], required=True)
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    parser.add_argument(
        "--output-root",
        "--output-dir",
        dest="output_root",
        type=Path,
        default=DEFAULT_OUTPUT_ROOT,
        help="Date-first result root. Images are saved to <root>/<YYYYMMDD>/Sector_Industry.",
    )
    parser.add_argument(
        "--material-root",
        type=Path,
        default=DEFAULT_MATERIAL_ROOT,
        help="Calculation-data root. JSON is saved to <root>/<YYYYMMDD>.",
    )
    parser.add_argument(
        "--as-of",
        type=parse_iso_date,
        default=None,
        help="Completed-close cutoff. Default: yesterday in local time.",
    )
    parser.add_argument("--market", choices=["ALL", "STOXX600", "TOPIX"], default="ALL")
    parser.add_argument("--top-n", type=int, default=8)
    parser.add_argument(
        "--min-coverage",
        type=float,
        default=0.80,
        help="Minimum valid constituent-return coverage required for each market.",
    )
    parser.add_argument(
        "--allow-mixed-as-of",
        action="store_true",
        help="Allow STOXX600 and TOPIX to resolve to different completed sessions.",
    )
    return parser.parse_args()


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)) and 20_000 < float(value) < 80_000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    return None


def number(value: Any) -> float:
    if isinstance(value, bool):
        return np.nan
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return np.nan


def clean_group(value: Any, fallback: str = "Unclassified") -> str:
    text = str(value or "").strip()
    return text if text else fallback


def load_metric_block(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    anchor_row: int,
    ticker_to_row: dict[str, int],
    member_count: int,
) -> tuple[list[date], np.ndarray]:
    label = ws.cell(anchor_row, 7).value
    if not isinstance(label, str) or " | " not in label:
        raise ValueError(f"Metric anchor not found at {ws.title}!G{anchor_row}: {label!r}")

    header = next(
        ws.iter_rows(
            min_row=anchor_row + 1,
            max_row=anchor_row + 1,
            min_col=7,
            max_col=ws.max_column,
            values_only=True,
        )
    )
    date_columns: list[tuple[int, date]] = []
    for offset, value in enumerate(header):
        parsed = as_date(value)
        if parsed is not None:
            date_columns.append((offset, parsed))
    if not date_columns:
        raise ValueError(f"No observation dates in {ws.title} metric block {label}")

    dates = [value for _, value in date_columns]
    matrix = np.full((member_count, len(dates)), np.nan, dtype=np.float64)
    rows = ws.iter_rows(
        min_row=anchor_row + 2,
        max_row=anchor_row + 1 + member_count,
        min_col=7,
        max_col=ws.max_column,
        values_only=True,
    )
    for row in rows:
        ticker = row[0]
        if not isinstance(ticker, str):
            continue
        out_row = ticker_to_row.get(ticker.strip())
        if out_row is None:
            continue
        for out_col, (offset, _) in enumerate(date_columns):
            if offset < len(row):
                matrix[out_row, out_col] = number(row[offset])
    return dates, matrix


def align_matrix(source_dates: list[date], source: np.ndarray, target_dates: list[date]) -> np.ndarray:
    if source_dates == target_dates:
        return source
    lookup = {value: idx for idx, value in enumerate(source_dates)}
    result = np.full((source.shape[0], len(target_dates)), np.nan, dtype=np.float64)
    for out_col, value in enumerate(target_dates):
        source_col = lookup.get(value)
        if source_col is not None:
            result[:, out_col] = source[:, source_col]
    return result


def load_panel_from_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    master: Path,
    config: dict[str, str],
    cutoff: date,
) -> MarketPanel:
    live_value = ws.cell(3, 6).value
    if not isinstance(live_value, (int, float)) or int(live_value) <= 0:
        raise ValueError(f"Invalid live member count in {ws.title}!F3: {live_value!r}")
    member_count = int(live_value)

    metadata = list(
        ws.iter_rows(
            min_row=8,
            max_row=7 + member_count,
            min_col=1,
            max_col=4,
            values_only=True,
        )
    )
    tickers = [str(row[0] or "").strip() for row in metadata]
    if any(not ticker for ticker in tickers):
        raise ValueError(f"Blank ticker in {ws.title} current membership block")
    if len(set(tickers)) != member_count:
        raise ValueError(f"Duplicate ticker in {ws.title} current membership block")
    names = [str(row[1] or row[0]).strip() for row in metadata]
    sectors = [clean_group(row[2]) for row in metadata]
    industries = [clean_group(row[3]) for row in metadata]

    ticker_to_row = {ticker: idx for idx, ticker in enumerate(tickers)}
    first_anchor = member_count + 10
    block_stride = member_count + 2
    price_anchor = first_anchor + METRICS.index("Price") * block_stride
    mcap_anchor = first_anchor + METRICS.index("Market Cap (Daily)") * block_stride
    price_dates, prices = load_metric_block(ws, price_anchor, ticker_to_row, member_count)
    mcap_dates, market_caps = load_metric_block(ws, mcap_anchor, ticker_to_row, member_count)
    market_caps = align_matrix(mcap_dates, market_caps, price_dates)

    keep = [idx for idx, value in enumerate(price_dates) if value <= cutoff]
    if not keep:
        raise ValueError(f"{config['title']}: no observations on or before {cutoff.isoformat()}")

    return MarketPanel(
        key=config["key"],
        title=config["title"],
        source_sheet=config["sheet"],
        source_code=str(ws.cell(4, 2).value or config["source_code"]).strip(),
        currency=str(ws.cell(4, 4).value or config["currency"]).strip(),
        source_workbook=str(master.resolve()),
        tickers=tickers,
        names=names,
        sectors=sectors,
        industries=industries,
        dates=[price_dates[idx] for idx in keep],
        prices=prices[:, keep],
        market_caps=market_caps[:, keep],
    )


def load_panels(master: Path, selected: list[str], cutoff: date) -> list[MarketPanel]:
    workbook = openpyxl.load_workbook(master, read_only=True, data_only=True)
    try:
        panels = []
        for key in selected:
            config = MARKETS[key]
            if config["sheet"] not in workbook.sheetnames:
                raise KeyError(f"Missing sheet: {config['sheet']}")
            panels.append(load_panel_from_sheet(workbook[config["sheet"]], master, config, cutoff))
        return panels
    finally:
        workbook.close()


def actual_sessions(prices: np.ndarray) -> list[int]:
    """Return genuine close columns, including the first usable close as the baseline.

    Bloomberg may repeat Friday's close on weekend/future-date columns. A new session
    therefore needs broad numeric coverage and price changes in more than 2% of valid
    pairs. Comparisons use the last accepted session, not merely the previous column.
    """
    member_count, column_count = prices.shape
    min_valid = max(10, math.ceil(member_count * 0.50))
    sessions: list[int] = []
    for current in range(column_count):
        current_values = prices[:, current]
        current_valid = np.isfinite(current_values) & (current_values > 0)
        if int(current_valid.sum()) < min_valid:
            continue
        if not sessions:
            sessions.append(current)
            continue
        before = prices[:, sessions[-1]]
        valid = current_valid & np.isfinite(before) & (before > 0)
        valid_count = int(valid.sum())
        if valid_count < min_valid:
            continue
        changed_count = int((~np.isclose(current_values[valid], before[valid], rtol=1e-7, atol=1e-9)).sum())
        if changed_count >= max(5, math.ceil(valid_count * 0.02)):
            sessions.append(current)
    return sessions


def safe_ratio(after: np.ndarray, before: np.ndarray) -> np.ndarray:
    result = np.full(before.shape, np.nan, dtype=np.float64)
    valid = np.isfinite(before) & np.isfinite(after) & (before > 0) & (after > 0)
    result[valid] = after[valid] / before[valid] - 1.0
    return result


def json_number(value: float | None) -> float | None:
    if value is None or not math.isfinite(float(value)):
        return None
    return float(value)


def corporate_action_reason(price_return: float, mcap_return: float | None) -> str | None:
    if price_return <= -0.90 or price_return >= 3.00:
        return "extreme_price_move"
    if mcap_return is None or 1.0 + price_return <= 0 or 1.0 + mcap_return <= 0:
        return None
    if abs(price_return) >= 0.35 and abs(math.log1p(price_return) - math.log1p(mcap_return)) >= 0.25:
        return "possible_corporate_action"
    return None


def summarize(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(records)
    values = [float(row["periodReturn"]) for row in rows if row.get("periodReturn") is not None and not row["excluded"]]
    return {
        "members": len(rows),
        "validMembers": len(values),
        "coverage": json_number(len(values) / len(rows)) if rows else None,
        "return": json_number(float(np.mean(values))) if values else None,
        "median": json_number(float(np.median(values))) if values else None,
        "breadth": json_number(sum(value > 0 for value in values) / len(values)) if values else None,
        "dispersion": json_number(float(np.std(values, ddof=0))) if values else None,
        "p25": json_number(float(np.percentile(values, 25))) if values else None,
        "p75": json_number(float(np.percentile(values, 75))) if values else None,
    }


def compact_stock(row: dict[str, Any]) -> dict[str, Any]:
    return {key: row.get(key) for key in ("ticker", "name", "industry", "periodReturn", "marketCapReturn")}


def stock_slice(records: list[dict[str, Any]], top_n: int) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    valid = [row for row in records if row.get("periodReturn") is not None and not row["excluded"]]
    ranked = sorted(valid, key=lambda row: float(row["periodReturn"]), reverse=True)
    top = [compact_stock(row) for row in ranked if float(row["periodReturn"]) > 0][:top_n]
    bottom = [compact_stock(row) for row in reversed(ranked) if float(row["periodReturn"]) < 0][:top_n]
    return top, bottom


def leadership_state(group_return: float | None, benchmark: float | None, breadth: float | None) -> str:
    if group_return is None or benchmark is None or breadth is None:
        return "판정 유보"
    if group_return >= benchmark and breadth >= 0.50:
        return "강세 확산"
    if group_return >= benchmark:
        return "상대 강세 집중"
    if breadth >= 0.50:
        return "방어적 확산"
    return "약세 확산"


def build_market(panel: MarketPanel, period: str, top_n: int, min_coverage: float) -> dict[str, Any]:
    span = 1 if period == "1D" else 5
    sessions = actual_sessions(panel.prices)
    required = span + 1
    if len(sessions) < required:
        found_dates = [panel.dates[idx].isoformat() for idx in sessions]
        raise ValueError(
            f"{panel.title}: {period} requires {required} genuine closes; found {len(sessions)} ({found_dates})"
        )
    window = sessions[-required:]
    start_idx, end_idx = window[0], window[-1]
    price_returns = safe_ratio(panel.prices[:, end_idx], panel.prices[:, start_idx])
    mcap_returns = safe_ratio(panel.market_caps[:, end_idx], panel.market_caps[:, start_idx])

    stocks: list[dict[str, Any]] = []
    for idx, ticker in enumerate(panel.tickers):
        price_return = json_number(price_returns[idx])
        mcap_return = json_number(mcap_returns[idx])
        reason = "missing_price" if price_return is None else corporate_action_reason(price_return, mcap_return)
        stocks.append(
            {
                "ticker": ticker,
                "name": panel.names[idx],
                "sector": panel.sectors[idx],
                "industry": panel.industries[idx],
                "startPrice": json_number(panel.prices[idx, start_idx]),
                "endPrice": json_number(panel.prices[idx, end_idx]),
                "periodReturn": price_return,
                "startMarketCap": json_number(panel.market_caps[idx, start_idx]),
                "endMarketCap": json_number(panel.market_caps[idx, end_idx]),
                "marketCapReturn": mcap_return,
                "excluded": reason is not None,
                "exclusionReason": reason,
            }
        )

    overall = summarize(stocks)
    if overall["coverage"] is None or overall["coverage"] < min_coverage:
        raise ValueError(
            f"{panel.title}: return coverage {overall['coverage']!r} is below {min_coverage:.0%}"
        )
    benchmark = overall["return"]
    sector_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    industry_rows: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in stocks:
        sector_rows[row["sector"]].append(row)
        industry_rows[(row["sector"], row["industry"])].append(row)

    industries: list[dict[str, Any]] = []
    for (sector, industry), rows in industry_rows.items():
        summary = summarize(rows)
        summary.update({"sector": sector, "industry": industry})
        summary["excessVsMarket"] = (
            json_number(summary["return"] - benchmark)
            if summary["return"] is not None and benchmark is not None
            else None
        )
        industries.append(summary)

    sectors: list[dict[str, Any]] = []
    for sector, rows in sector_rows.items():
        summary = summarize(rows)
        sector_industries = [
            row for row in industries if row["sector"] == sector and row["return"] is not None
        ]
        sector_industries.sort(key=lambda row: float(row["return"]), reverse=True)
        top_gainers, bottom_losers = stock_slice(rows, top_n)
        summary.update(
            {
                "sector": sector,
                "industryCount": len(sector_industries),
                "excessVsMarket": (
                    json_number(summary["return"] - benchmark)
                    if summary["return"] is not None and benchmark is not None
                    else None
                ),
                "state": leadership_state(summary["return"], benchmark, summary["breadth"]),
                "topIndustry": sector_industries[0] if sector_industries else None,
                "bottomIndustry": sector_industries[-1] if sector_industries else None,
                "industrySpread": (
                    json_number(sector_industries[0]["return"] - sector_industries[-1]["return"])
                    if len(sector_industries) >= 2
                    else 0.0 if len(sector_industries) == 1 else None
                ),
                "topGainers": top_gainers,
                "bottomLosers": bottom_losers,
            }
        )
        sectors.append(summary)

    sectors.sort(key=lambda row: float(row["return"]) if row["return"] is not None else -math.inf, reverse=True)
    industries.sort(key=lambda row: float(row["return"]) if row["return"] is not None else -math.inf, reverse=True)
    for rank, row in enumerate(sectors, 1):
        row["rank"] = rank

    return {
        "key": panel.key,
        "title": panel.title,
        "sourceSheet": panel.source_sheet,
        "sourceIndexLabel": panel.source_code,
        "currency": panel.currency,
        "sourceWorkbook": panel.source_workbook,
        "period": period,
        "asOf": panel.dates[end_idx].isoformat(),
        "startDate": panel.dates[start_idx].isoformat(),
        "sessionDates": [panel.dates[idx].isoformat() for idx in window],
        "memberCount": len(stocks),
        "validMemberCount": overall["validMembers"],
        "excludedCount": sum(row["excluded"] for row in stocks),
        "overall": overall,
        "sectors": sectors,
        "industries": industries,
        "stocks": stocks,
        "tableImage": None,
    }


def validate_market(market: dict[str, Any], top_n: int) -> None:
    if market["memberCount"] != len(market["stocks"]):
        raise AssertionError(f"{market['key']}: member count mismatch")
    if len({row["ticker"] for row in market["stocks"]}) != market["memberCount"]:
        raise AssertionError(f"{market['key']}: duplicate ticker")
    if sum(row["members"] for row in market["sectors"]) != market["memberCount"]:
        raise AssertionError(f"{market['key']}: sector hierarchy mismatch")
    if sum(row["members"] for row in market["industries"]) != market["memberCount"]:
        raise AssertionError(f"{market['key']}: industry hierarchy mismatch")
    required = 2 if market["period"] == "1D" else 6
    if len(market["sessionDates"]) != required:
        raise AssertionError(f"{market['key']}: session window mismatch")
    for sector in market["sectors"]:
        if len(sector["topGainers"]) > top_n or len(sector["bottomLosers"]) > top_n:
            raise AssertionError(f"{market['key']} {sector['sector']}: top/bottom limit exceeded")
        top = {row["ticker"] for row in sector["topGainers"]}
        bottom = {row["ticker"] for row in sector["bottomLosers"]}
        if top & bottom:
            raise AssertionError(f"{market['key']} {sector['sector']}: top/bottom overlap")
        if any(row["periodReturn"] <= 0 for row in sector["topGainers"]):
            raise AssertionError(f"{market['key']} {sector['sector']}: non-positive gainer")
        if any(row["periodReturn"] >= 0 for row in sector["bottomLosers"]):
            raise AssertionError(f"{market['key']} {sector['sector']}: non-negative loser")


def font(size: int, bold: bool = False, serif: bool = False) -> ImageFont.ImageFont:
    if serif:
        candidates = [
            Path("C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf"),
            Path("C:/Windows/Fonts/malgunbd.ttf" if bold else "C:/Windows/Fonts/malgun.ttf"),
        ]
    else:
        candidates = [
            Path("C:/Windows/Fonts/malgunbd.ttf" if bold else "C:/Windows/Fonts/malgun.ttf"),
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        ]
    found = next((path for path in candidates if path.exists()), None)
    return ImageFont.truetype(str(found), size=size) if found else ImageFont.load_default()


def percent(value: float | None, decimals: int = 1, points: bool = False) -> str:
    if value is None:
        return "—"
    suffix = "%p" if points else "%"
    scaled = round(value * 100, decimals)
    if scaled == 0:
        return f"{0:.{decimals}f}{suffix}"
    return f"{scaled:+.{decimals}f}{suffix}"


def breadth(value: float | None) -> str:
    return "—" if value is None else f"{value * 100:.0f}%"


def text_color(value: float | None) -> str:
    # The table displays one decimal percentage point, so a rounded zero should
    # also be neutral rather than a red/green '-0.0'.
    if value is None or round(value * 100, 1) == 0:
        return COLORS["black"]
    return COLORS["green"] if value > 0 else COLORS["red"]


def centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, text_font: ImageFont.ImageFont, fill: str) -> None:
    left, top, right, bottom = box
    bbox = draw.textbbox((0, 0), text, font=text_font)
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.text(
        (left + (right - left - width) / 2, top + (bottom - top - height) / 2 - bbox[1]),
        text,
        font=text_font,
        fill=fill,
    )


def render_table(market: dict[str, Any], output: Path) -> None:
    columns = [
        ("섹터", 420),
        ("N", 115),
        ("산업그룹", 145),
        (f"{market['period']} 평균", 185),
        ("중앙값", 185),
        ("상승 비율", 180),
        ("분산", 180),
        ("시장 대비", 190),
    ]
    width = sum(value for _, value in columns) + 2
    header_height = 68
    row_height = 62
    height = header_height + row_height * len(market["sectors"]) + 2
    image = Image.new("RGB", (width, height), COLORS["white"])
    draw = ImageDraw.Draw(image)
    header_font = font(25, bold=True)
    sector_font = font(25, bold=True, serif=True)
    number_font = font(25, serif=True)

    x_positions = [1]
    for _, column_width in columns:
        x_positions.append(x_positions[-1] + column_width)

    draw.rectangle((1, 1, width - 2, header_height), fill=COLORS["navy"])
    for idx, (label, _) in enumerate(columns):
        centered_text(draw, (x_positions[idx], 1, x_positions[idx + 1], header_height), label, header_font, COLORS["white"])

    for row_idx, sector in enumerate(market["sectors"]):
        top = header_height + row_idx * row_height
        bottom = top + row_height
        if row_idx % 2:
            draw.rectangle((1, top, width - 2, bottom), fill=COLORS["alternate"])
        values = [
            sector["sector"],
            str(sector["validMembers"]),
            str(sector["industryCount"]),
            percent(sector["return"]),
            percent(sector["median"]),
            breadth(sector["breadth"]),
            percent(sector["dispersion"]),
            percent(sector["excessVsMarket"], points=True),
        ]
        colors = [
            COLORS["black"],
            COLORS["black"],
            COLORS["black"],
            text_color(sector["return"]),
            text_color(sector["median"]),
            COLORS["black"],
            COLORS["black"],
            text_color(sector["excessVsMarket"]),
        ]
        for col_idx, value in enumerate(values):
            box = (x_positions[col_idx], top, x_positions[col_idx + 1], bottom)
            if col_idx == 0:
                bbox = draw.textbbox((0, 0), value, font=sector_font)
                text_height = bbox[3] - bbox[1]
                draw.text((box[0] + 16, top + (row_height - text_height) / 2 - bbox[1]), value, font=sector_font, fill=colors[col_idx])
            else:
                centered_text(draw, box, value, number_font, colors[col_idx])

    for x in x_positions:
        draw.line((x, 1, x, height - 2), fill=COLORS["grid"], width=2)
    draw.line((1, 1, width - 2, 1), fill=COLORS["grid"], width=2)
    draw.line((1, header_height, width - 2, header_height), fill=COLORS["grid"], width=2)
    for row_idx in range(1, len(market["sectors"]) + 1):
        y = header_height + row_idx * row_height
        draw.line((1, y, width - 2, y), fill=COLORS["grid"], width=2)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    image.save(temporary, format="PNG", optimize=True)
    os.replace(temporary, output)


def output_tag(markets: list[dict[str, Any]], allow_mixed: bool) -> str:
    dates = sorted({market["asOf"] for market in markets})
    if len(dates) == 1:
        return dates[0].replace("-", "")
    detail = ", ".join(f"{market['key']}={market['asOf']}" for market in markets)
    if not allow_mixed:
        raise ValueError(
            f"Mixed market as-of dates are not allowed: {detail}. Use --allow-mixed-as-of only if intentional."
        )
    return f"MIXED_{dates[0].replace('-', '')}_{dates[-1].replace('-', '')}"


def atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def main() -> int:
    args = parse_args()
    if not 1 <= args.top_n <= 8:
        raise ValueError("--top-n must be between 1 and 8")
    if not 0.50 <= args.min_coverage <= 1.0:
        raise ValueError("--min-coverage must be between 0.50 and 1.00")

    master = args.master.resolve()
    if not master.exists():
        raise FileNotFoundError(master)
    cutoff = args.as_of or (date.today() - timedelta(days=1))
    selected = list(MARKETS) if args.market == "ALL" else [args.market]
    panels = load_panels(master, selected, cutoff)
    markets = [build_market(panel, args.period, args.top_n, args.min_coverage) for panel in panels]
    for market in markets:
        validate_market(market, args.top_n)

    computed_tag = output_tag(markets, args.allow_mixed_as_of)
    # Keep the storage hierarchy strictly date-first even when two exchanges
    # resolve to different genuine closes. Market-specific dates remain in the
    # filenames and payload; the requested cutoff is the shared batch date.
    tag = cutoff.strftime("%Y%m%d") if computed_tag.startswith("MIXED_") else computed_tag
    output_run_dir = args.output_root.resolve() / tag / "Sector_Industry"
    material_run_dir = args.material_root.resolve() / tag
    for category in ("AI", "Sector_Industry", "Top10_Down10"):
        (args.output_root.resolve() / tag / category).mkdir(parents=True, exist_ok=True)
    for market in markets:
        market_tag = market["asOf"].replace("-", "")
        image_path = output_run_dir / f"{market['key']}_섹터_{args.period}_{market_tag}.png"
        render_table(market, image_path)
        market["tableImage"] = str(image_path.resolve())

    payload = {
        "meta": {
            "schemaVersion": 2,
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "requestedCutoff": cutoff.isoformat(),
            "storageDate": tag,
            "period": args.period,
            "periodDefinition": (
                "latest completed close divided by the previous genuine close minus one"
                if args.period == "1D"
                else "latest completed close divided by the close five genuine trading sessions earlier minus one"
            ),
            "sessionDetection": "first usable close plus later columns with >=50% pair coverage and price changes in >=max(5 stocks, 2% of valid pairs)",
            "aggregation": "equal-weight arithmetic mean of valid constituent local-price returns",
            "hierarchy": "GICS Sector > GICS Industry Group > constituent",
            "topBottomLimit": args.top_n,
            "minimumCoverage": args.min_coverage,
            "notes": [
                "Returns exclude dividends and FX translation.",
                "Current constituents are backfilled through the observation window, so survivorship bias remains.",
                "Obvious corporate-action-like price jumps are excluded when price and market-cap continuity diverge.",
                "Future-date, weekend, holiday, and unchanged carry columns are not counted as genuine sessions.",
            ],
        },
        "markets": markets,
    }
    market_tag = "_".join(market["key"] for market in markets)
    json_path = material_run_dir / f"{market_tag}_Sector_Industry_{args.period}_{tag}.json"
    atomic_json(json_path, payload)
    latest_pointer = args.material_root.resolve() / f"latest_{args.period}.json"
    atomic_json(
        latest_pointer,
        {
            "period": args.period,
            "asOf": {market["key"]: market["asOf"] for market in markets},
            "material": str(json_path.resolve()),
            "json": str(json_path.resolve()),
            "images": {market["key"]: market["tableImage"] for market in markets},
            "updatedAt": payload["meta"]["generatedAt"],
        },
    )
    print(json_path)
    for market in markets:
        print(market["tableImage"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
