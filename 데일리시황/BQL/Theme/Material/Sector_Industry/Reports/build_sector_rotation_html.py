from __future__ import annotations

import argparse
import hashlib
import html
import json
import math
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl


HERE = Path(__file__).resolve().parent
DEFAULT_BQL_ROOT = Path(r"C:\Users\infomax\Documents\python\QAE\데일리시황\BQL")
DEFAULT_MASTER = DEFAULT_BQL_ROOT / "Rawfile" / "BQuant_Master.xlsx"
DEFAULT_OUTPUT_ROOT = DEFAULT_BQL_ROOT / "Theme" / "output"
DEFAULT_DESCRIPTION_BOOK = DEFAULT_BQL_ROOT / "Theme" / "Material" / "Sector_Industry" / "stock_descriptions.xlsx"

MARKETS = {
    "STOXX600": {
        "title": "STOXX Europe 600",
        "sheet": "SXXP Raw",
        "code": "SXXP Index",
    },
    "SP500": {
        "title": "S&P 500",
        "sheet": "SPX Raw",
        "code": "SPX Index",
    },
    "TOPIX": {
        "title": "TOPIX",
        "sheet": "TPX500 Raw",
        "code": "TPX500 Index",
    },
    "HANGSENG": {
        "title": "Hang Seng",
        "sheet": "HSI Raw",
        "code": "HSI Index",
    },
}

PERIODS = ("1D", "5D", "1M", "3M", "6M")
TRADING_INTERVALS = {"1D": 1, "5D": 5}
CALENDAR_MONTHS = {"1M": 1, "3M": 3, "6M": 6}


@dataclass
class MarketRaw:
    key: str
    title: str
    source_sheet: str
    source_code: str
    tickers: list[str]
    names: list[str]
    sectors: list[str]
    industries: list[str]
    dates: list[date]
    prices: np.ndarray
    market_caps: np.ndarray


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate fast HTML sector/industry rotation reports for Europe, US, Japan, and Hong Kong."
    )
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--primary", choices=("1D", "5D"), required=True)
    parser.add_argument(
        "--market",
        choices=("ALL", "STOXX600", "SP500", "TOPIX", "HANGSENG"),
        default="ALL",
    )
    parser.add_argument("--top-n", type=int, default=5)
    parser.add_argument("--as-of", type=date.fromisoformat)
    parser.add_argument("--descriptions", type=Path, default=DEFAULT_DESCRIPTION_BOOK)
    return parser.parse_args()


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        if 20_000 < float(value) < 80_000:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    return None


def number(value: Any) -> float:
    if isinstance(value, bool):
        return np.nan
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return np.nan


def clean_group(value: Any) -> str:
    text = str(value or "").strip()
    return text if text else "Unclassified"


def load_descriptions(path: Path) -> dict[tuple[str, str], str]:
    if not path.exists():
        return {}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = "Stock_Descriptions"
        if sheet_name not in workbook.sheetnames:
            return {}
        ws = workbook[sheet_name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = {str(value).strip(): idx for idx, value in enumerate(header) if value is not None}
        required = ("Market", "Ticker", "Business Summary (KR)")
        if any(name not in columns for name in required):
            return {}
        result: dict[tuple[str, str], str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            market = str(row[columns["Market"]] or "").strip()
            ticker = str(row[columns["Ticker"]] or "").strip()
            summary = str(row[columns["Business Summary (KR)"]] or "").strip()
            if market and ticker and summary:
                result[(market, ticker)] = summary
        return result
    finally:
        workbook.close()


def metric_anchor(member_count: int, metric_index: int) -> int:
    first_anchor = member_count + 10
    block_stride = member_count + 2
    return first_anchor + metric_index * block_stride


def load_price_block(ws: Any, anchor_row: int, ticker_to_row: dict[str, int], n: int) -> tuple[list[date], np.ndarray]:
    label = ws.cell(anchor_row, 7).value
    if not isinstance(label, str) or " | " not in label:
        raise ValueError(f"가격 블록을 찾지 못했습니다: {ws.title}!G{anchor_row} ({label!r})")

    header = next(
        ws.iter_rows(
            min_row=anchor_row + 1,
            max_row=anchor_row + 1,
            min_col=7,
            max_col=ws.max_column,
            values_only=True,
        )
    )
    date_cols: list[tuple[int, date]] = []
    for offset, value in enumerate(header):
        parsed = as_date(value)
        if parsed is not None:
            date_cols.append((offset, parsed))
    if not date_cols:
        raise ValueError(f"가격 날짜가 없습니다: {ws.title}")

    dates = [obs_date for _, obs_date in date_cols]
    matrix = np.full((n, len(dates)), np.nan, dtype=np.float64)
    rows = ws.iter_rows(
        min_row=anchor_row + 2,
        max_row=anchor_row + 1 + n,
        min_col=7,
        max_col=ws.max_column,
        values_only=True,
    )
    for row in rows:
        ticker = str(row[0] or "").strip()
        out_row = ticker_to_row.get(ticker)
        if out_row is None:
            continue
        for out_col, (offset, _) in enumerate(date_cols):
            if offset < len(row):
                matrix[out_row, out_col] = number(row[offset])
    return dates, matrix


def load_market(workbook: Any, key: str) -> MarketRaw:
    config = MARKETS[key]
    ws = workbook[config["sheet"]]
    live_value = ws.cell(3, 6).value
    if not isinstance(live_value, (int, float)) or int(live_value) <= 0:
        raise ValueError(f"구성종목 수가 올바르지 않습니다: {config['sheet']}!F3")
    member_count = int(live_value)

    tickers: list[str] = []
    names: list[str] = []
    sectors: list[str] = []
    industries: list[str] = []
    for row in ws.iter_rows(min_row=8, max_row=7 + member_count, min_col=1, max_col=4, values_only=True):
        ticker = str(row[0] or "").strip()
        if not ticker:
            continue
        tickers.append(ticker)
        names.append(str(row[1] or ticker).strip())
        sectors.append(clean_group(row[2]))
        industries.append(clean_group(row[3]))
    if len(tickers) != member_count:
        raise ValueError(f"{key}: 구성종목 {member_count}개 중 메타데이터는 {len(tickers)}개입니다.")
    if len(set(tickers)) != len(tickers):
        raise ValueError(f"{key}: 구성종목 티커가 중복됩니다.")

    ticker_to_row = {ticker: idx for idx, ticker in enumerate(tickers)}
    # Price and daily market cap are the fifth and sixth metric blocks.
    dates, prices = load_price_block(ws, metric_anchor(member_count, 4), ticker_to_row, member_count)
    cap_dates, market_caps = load_price_block(ws, metric_anchor(member_count, 5), ticker_to_row, member_count)
    if cap_dates != dates:
        raise ValueError(f"{key}: 가격과 일별 시가총액의 날짜가 일치하지 않습니다.")
    return MarketRaw(
        key=key,
        title=config["title"],
        source_sheet=config["sheet"],
        source_code=str(ws.cell(4, 2).value or config["code"]).strip(),
        tickers=tickers,
        names=names,
        sectors=sectors,
        industries=industries,
        dates=dates,
        prices=prices,
        market_caps=market_caps,
    )


def genuine_sessions(prices: np.ndarray) -> list[int]:
    accepted: list[int] = []
    for current in range(1, prices.shape[1]):
        before = prices[:, current - 1]
        after = prices[:, current]
        valid = np.isfinite(before) & np.isfinite(after) & (before > 0) & (after > 0)
        comparable = int(valid.sum())
        if comparable < max(10, int(prices.shape[0] * 0.50)):
            continue
        changed = ~np.isclose(after[valid], before[valid], rtol=1e-7, atol=1e-9)
        if int(changed.sum()) >= max(5, math.ceil(comparable * 0.02)):
            accepted.append(current)
    # The first accepted change also proves that the immediately preceding column
    # is a genuine usable closing baseline. Keeping it is essential for an exact
    # five-session return when only six closing observations have accumulated.
    if accepted:
        baseline = accepted[0] - 1
        base_values = prices[:, baseline]
        valid_base = np.isfinite(base_values) & (base_values > 0)
        if int(valid_base.sum()) >= max(10, int(prices.shape[0] * 0.50)):
            accepted.insert(0, baseline)
    return accepted


def shift_months(value: date, months: int) -> date:
    year = value.year
    month = value.month - months
    while month <= 0:
        year -= 1
        month += 12
    month_lengths = (31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31)
    return date(year, month, min(value.day, month_lengths[month - 1]))


def resolve_windows(raw: MarketRaw, cutoff: date | None) -> tuple[int, dict[str, int | None], list[int]]:
    sessions = genuine_sessions(raw.prices)
    if cutoff is not None:
        sessions = [idx for idx in sessions if raw.dates[idx] <= cutoff]
    if not sessions:
        raise ValueError(f"{raw.title}: 유효 거래일을 찾지 못했습니다.")
    end_idx = sessions[-1]
    starts: dict[str, int | None] = {}
    position = len(sessions) - 1
    for period, intervals in TRADING_INTERVALS.items():
        target_position = position - intervals
        starts[period] = sessions[target_position] if target_position >= 0 else None
    for period, months in CALENDAR_MONTHS.items():
        target_date = shift_months(raw.dates[end_idx], months)
        candidates = [idx for idx in sessions if raw.dates[idx] <= target_date]
        starts[period] = candidates[-1] if candidates else None
    return end_idx, starts, sessions


def safe_returns(after: np.ndarray, before: np.ndarray) -> np.ndarray:
    result = np.full(before.shape, np.nan, dtype=np.float64)
    valid = np.isfinite(before) & np.isfinite(after) & (before > 0) & (after > 0)
    result[valid] = after[valid] / before[valid] - 1.0
    # Remove obvious split/corporate-action-like observations from return summaries.
    result[(result <= -0.90) | (result >= 3.00)] = np.nan
    return result


def summarize_group(
    stock_indices: list[int],
    returns: dict[str, np.ndarray],
    start_caps: dict[str, np.ndarray],
) -> dict[str, Any]:
    summary: dict[str, Any] = {"n": len(stock_indices), "returns": {}, "breadth": {}}
    for period in PERIODS:
        vector = returns.get(period)
        if vector is None:
            summary["returns"][period] = None
            summary["breadth"][period] = None
            continue
        values = [float(vector[idx]) for idx in stock_indices if math.isfinite(float(vector[idx]))]
        caps = start_caps.get(period)
        if caps is None:
            summary["returns"][period] = None
        else:
            eligible = [idx for idx in stock_indices if math.isfinite(float(vector[idx])) and math.isfinite(float(caps[idx])) and float(caps[idx]) > 0]
            summary["returns"][period] = (
                float(np.average(vector[eligible], weights=caps[eligible])) if eligible else None
            )
        summary["breadth"][period] = (sum(value > 0 for value in values) / len(values)) if values else None
    return summary


def build_market(raw: MarketRaw, primary: str, top_n: int, cutoff: date | None) -> dict[str, Any]:
    end_idx, starts, sessions = resolve_windows(raw, cutoff)
    returns: dict[str, np.ndarray] = {}
    start_caps: dict[str, np.ndarray] = {}
    for period, start_idx in starts.items():
        if start_idx is not None:
            returns[period] = safe_returns(raw.prices[:, end_idx], raw.prices[:, start_idx])
            start_caps[period] = raw.market_caps[:, start_idx]

    all_indices = list(range(len(raw.tickers)))
    benchmark = summarize_group(all_indices, returns, start_caps)
    sector_map: dict[str, list[int]] = {}
    industry_map: dict[tuple[str, str], list[int]] = {}
    for idx, (sector, industry) in enumerate(zip(raw.sectors, raw.industries)):
        sector_map.setdefault(sector, []).append(idx)
        industry_map.setdefault((sector, industry), []).append(idx)

    industries: list[dict[str, Any]] = []
    for (sector, industry), indices in industry_map.items():
        row = summarize_group(indices, returns, start_caps)
        row.update({"sector": sector, "industry": industry, "indices": indices})
        industries.append(row)

    sectors: list[dict[str, Any]] = []
    for sector, indices in sector_map.items():
        row = summarize_group(indices, returns, start_caps)
        ranked_indices: list[int] = []
        primary_vector = returns.get(primary)
        if primary_vector is not None:
            ranked_indices = [idx for idx in indices if math.isfinite(float(primary_vector[idx]))]
            ranked_indices.sort(key=lambda idx: float(primary_vector[idx]), reverse=True)
        top = ranked_indices[:top_n]
        bottom = list(reversed(ranked_indices[-top_n:]))
        row.update(
            {
                "sector": sector,
                "indices": indices,
                "industries": [item for item in industries if item["sector"] == sector],
                "top": top,
                "bottom": bottom,
            }
        )
        sectors.append(row)

    def sort_value(row: dict[str, Any]) -> float:
        value = row["returns"].get(primary)
        return float(value) if value is not None else -math.inf

    sectors.sort(key=sort_value, reverse=True)
    industries.sort(key=sort_value, reverse=True)
    for sector in sectors:
        sector["industries"].sort(key=sort_value, reverse=True)

    coverage: dict[str, float | None] = {}
    cap_coverage: dict[str, float | None] = {}
    for period in PERIODS:
        vector = returns.get(period)
        coverage[period] = float(np.isfinite(vector).sum() / len(vector)) if vector is not None else None
        caps = start_caps.get(period)
        cap_coverage[period] = float((np.isfinite(caps) & (caps > 0) & np.isfinite(vector)).sum() / len(vector)) if caps is not None and vector is not None else None
    if coverage.get(primary) is None:
        raise ValueError(f"{raw.title}: {primary} 계산에 필요한 거래일 이력이 부족합니다.")
    if float(coverage[primary]) < 0.80:
        raise ValueError(
            f"{raw.title}: {primary} 유효 가격 비율이 {float(coverage[primary]) * 100:.1f}%로 80% 미만입니다."
        )
    if cap_coverage[primary] is None or float(cap_coverage[primary]) < 0.80:
        raise ValueError(f"{raw.title}: {primary} 유효 시가총액·가격 비율이 80% 미만입니다.")

    return {
        "key": raw.key,
        "title": raw.title,
        "sourceSheet": raw.source_sheet,
        "sourceCode": raw.source_code,
        "asOf": raw.dates[end_idx].isoformat(),
        "primary": primary,
        "periodStart": {period: raw.dates[idx].isoformat() if idx is not None else None for period, idx in starts.items()},
        "memberCount": len(raw.tickers),
        "coverage": coverage,
        "capCoverage": cap_coverage,
        "benchmark": benchmark,
        "sectors": sectors,
        "industries": industries,
        "stock": {
            "ticker": raw.tickers,
            "name": raw.names,
            "sector": raw.sectors,
            "industry": raw.industries,
            "returns": {period: vector.tolist() for period, vector in returns.items()},
        },
        "sessionDates": [raw.dates[idx].isoformat() for idx in sessions[-10:]],
    }


def esc(value: Any) -> str:
    return html.escape(str(value), quote=True)


def commentary_slot(kind: str, *parts: Any) -> str:
    """Return a stable, invisible hook for prose-only post processing."""
    raw = "\x1f".join([kind, *(str(part) for part in parts)])
    slot_id = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:20]
    return (
        f'<div class="commentary-slot" data-commentary-slot="{esc(kind)}" '
        f'data-commentary-id="{slot_id}"></div>'
    )


def pct(value: float | None, digits: int = 1) -> str:
    if value is None or not math.isfinite(float(value)):
        return "데이터 부족"
    return f"{float(value) * 100:+.{digits}f}%"


def breadth_pct(value: float | None) -> str:
    if value is None or not math.isfinite(float(value)):
        return "데이터 부족"
    return f"{float(value) * 100:.0f}%"


def value_class(value: float | None) -> str:
    if value is None or not math.isfinite(float(value)):
        return "na"
    return "pos" if value > 0 else "neg" if value < 0 else "zero"


def return_cells(values: dict[str, float | None]) -> str:
    return "".join(
        f'<td class="num {value_class(values.get(period))}">{pct(values.get(period))}</td>'
        for period in PERIODS
    )


def multi_period_table(rows: list[dict[str, Any]], name_columns: list[tuple[str, str]], benchmark: dict[str, Any] | None = None) -> str:
    head = "".join(f"<th>{esc(label)}</th>" for _, label in name_columns)
    body: list[str] = []
    for row in rows:
        names = "".join(f"<td>{esc(row.get(key, ''))}</td>" for key, _ in name_columns)
        body.append(f"<tr>{names}{return_cells(row['returns'])}</tr>")
    if benchmark is not None:
        blanks = "".join("<td></td>" for _ in name_columns[1:])
        body.append(
            f"<tr class='benchmark'><td>BM</td>{blanks}{return_cells(benchmark['returns'])}</tr>"
        )
    return (
        "<div class='table-wrap'><table><thead><tr>"
        f"{head}"
        + "".join(f"<th class='num'>{period}</th>" for period in PERIODS)
        + "</tr></thead><tbody>"
        + "".join(body)
        + "</tbody></table></div>"
    )


def overall_table(market: dict[str, Any]) -> str:
    primary = market["primary"]
    bm_return = market["benchmark"]["returns"].get(primary)
    rows: list[str] = []
    for sector in market["sectors"]:
        value = sector["returns"].get(primary)
        excess = value - bm_return if value is not None and bm_return is not None else None
        rows.append(
            "<tr>"
            f"<td>{esc(sector['sector'])}</td><td class='num'>{sector['n']}</td>"
            f"<td class='num {value_class(value)}'>{pct(value)}</td>"
            f"<td class='num'>{breadth_pct(sector['breadth'].get(primary))}</td>"
            f"<td class='num {value_class(excess)}'>{pct(excess)}</td>"
            "</tr>"
        )
    bm_breadth = market["benchmark"]["breadth"].get(primary)
    rows.append(
        "<tr class='benchmark'>"
        f"<td>BM 전체 (시작일 시총가중)</td><td class='num'>{market['memberCount']}</td>"
        f"<td class='num {value_class(bm_return)}'>{pct(bm_return)}</td>"
        f"<td class='num'>{breadth_pct(bm_breadth)}</td>"
        "<td class='num zero'>0.0%</td>"
        "</tr>"
    )
    return (
        "<div class='table-wrap'><table><thead><tr>"
        f"<th>섹터</th><th class='num'>N</th><th class='num'>{primary} 시총가중</th><th class='num'>{primary} 상승 종목 비율</th><th class='num'>BM 대비</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></div>"
    )


def breadth_table(market: dict[str, Any]) -> str:
    primary = market["primary"]
    rows = []
    for sector in market["sectors"]:
        value = sector["breadth"].get(primary)
        rows.append(f"<tr><td>{esc(sector['sector'])}</td><td class='num'>{breadth_pct(value)}</td></tr>")
    bm = market["benchmark"]["breadth"].get(primary)
    rows.append(f"<tr class='benchmark'><td>BM 전체</td><td class='num'>{breadth_pct(bm)}</td></tr>")
    return (
        "<div class='table-wrap compact'><table><thead><tr>"
        f"<th>섹터</th><th class='num'>{primary} 상승 종목 비율</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></div>"
    )


def stock_table(market: dict[str, Any], sector: dict[str, Any], descriptions: dict[tuple[str, str], str]) -> str:
    primary = market["primary"]
    vectors = market["stock"]["returns"]
    primary_vector = vectors.get(primary, [])
    top = sector["top"]
    bottom = sector["bottom"]
    rows: list[str] = []
    for rank in range(max(len(top), len(bottom))):
        left = top[rank] if rank < len(top) else None
        right = bottom[rank] if rank < len(bottom) else None
        if left is None:
            left_cells = "<td></td><td></td>"
        else:
            left_value = primary_vector[left]
            left_ticker = market['stock']['ticker'][left]
            left_description = descriptions.get(
                (market["key"], left_ticker),
                f"{market['stock']['industry'][left]} · {market['stock']['sector'][left]}",
            )
            left_cells = (
                f"<td><div class='stock-name'>{esc(market['stock']['name'][left])}</div><div class='stock-desc'>{esc(left_description)}</div>"
                + commentary_slot("stock", market["key"], left_ticker)
                + "</td>"
                + f"<td class='num {value_class(left_value)}'>{pct(left_value)}</td>"
            )
        if right is None:
            right_cells = "<td></td><td></td>"
        else:
            right_value = primary_vector[right]
            right_ticker = market['stock']['ticker'][right]
            right_description = descriptions.get(
                (market["key"], right_ticker),
                f"{market['stock']['industry'][right]} · {market['stock']['sector'][right]}",
            )
            right_cells = (
                f"<td><div class='stock-name'>{esc(market['stock']['name'][right])}</div><div class='stock-desc'>{esc(right_description)}</div>"
                + commentary_slot("stock", market["key"], right_ticker)
                + "</td>"
                + f"<td class='num {value_class(right_value)}'>{pct(right_value)}</td>"
            )
        rows.append(f"<tr><td class='num'>{rank + 1}</td>{left_cells}<td class='num'>{rank + 1}</td>{right_cells}</tr>")
    return (
        "<div class='table-wrap'><table><thead><tr>"
        f"<th class='num'>#</th><th>상위 종목</th><th class='num'>{primary}</th>"
        f"<th class='num'>#</th><th>하위 종목</th><th class='num'>{primary}</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></div>"
    )


def key_findings(market: dict[str, Any]) -> str:
    primary = market["primary"]
    available = [row for row in market["industries"] if row["returns"].get(primary) is not None]
    leaders = available[:5]
    laggards = list(reversed(available[-5:]))
    sector_leader = market["sectors"][0] if market["sectors"] else None
    sector_laggard = market["sectors"][-1] if market["sectors"] else None
    cards = []
    if sector_leader:
        cards.append(f"<div class='card'><span>섹터 선두</span><b>{esc(sector_leader['sector'])}</b><strong class='pos'>{pct(sector_leader['returns'].get(primary))}</strong></div>")
    if sector_laggard:
        cards.append(f"<div class='card'><span>섹터 최하위</span><b>{esc(sector_laggard['sector'])}</b><strong class='neg'>{pct(sector_laggard['returns'].get(primary))}</strong></div>")
    cards.append(f"<div class='card'><span>BM 수익률</span><b>{primary}</b><strong class='{value_class(market['benchmark']['returns'].get(primary))}'>{pct(market['benchmark']['returns'].get(primary))}</strong></div>")
    selected = leaders + [row for row in laggards if row not in leaders]
    return "<div class='cards'>" + "".join(cards) + "</div>" + multi_period_table(selected, [("industry", "산업그룹"), ("sector", "섹터")])


def render_market_panel(market: dict[str, Any], descriptions: dict[tuple[str, str], str], active: bool = False) -> str:
    primary = market["primary"]
    starts = " · ".join(
        f"{period}: {market['periodStart'].get(period) or '데이터 부족'}"
        for period in PERIODS
    )
    sector_rotation = multi_period_table(market["sectors"], [("sector", "섹터")], market["benchmark"])
    industry_all = multi_period_table(market["industries"], [("industry", "산업그룹"), ("sector", "섹터")])
    details: list[str] = []
    for sector in market["sectors"]:
        details.append(
            f"<section class='sector-detail'><h3>{esc(sector['sector'])}</h3>"
            + commentary_slot("sector", market["key"], sector["sector"])
            + "<h4>산업그룹 성과</h4>"
            + multi_period_table(sector["industries"], [("industry", "산업그룹")])
            + f"<h4>{primary} 기준 상위·하위 종목 각 5개</h4>"
            + stock_table(market, sector, descriptions)
            + "</section>"
        )

    active_class = " active" if active else ""
    return f"""<section id="panel-{esc(market['key'])}" class="market-panel{active_class}" data-market="{esc(market['key'])}">
<div><span class="badge">{primary}</span><span class="badge">{esc(market['key'])}</span></div>
<h1>{esc(market['title'])} 섹터·산업 로테이션</h1>
<div class="meta">기준일 {esc(market['asOf'])} · 구성종목 {market['memberCount']}개 · 원자료 {esc(market['sourceSheet'])}<br>기간 시작일 · {esc(starts)}</div>
{commentary_slot("market", market["key"])}

<h2>1. 전체 섹터 성과표 — {primary}</h2>
{overall_table(market)}

<h2>2. 섹터 로테이션 — 1D / 5D / 1M / 3M / 6M</h2>
{sector_rotation}

<h2>3. 섹터별 상승 종목 비율 — {primary}</h2>
{breadth_table(market)}

<h2>4. Key Findings — 산업그룹 상·하위</h2>
{key_findings(market)}

<h2>5. 전체 산업그룹 성과표</h2>
{industry_all}

<h2>6. 섹터별 산업그룹과 종목</h2>
{''.join(details)}

<div class="note">계산 기준: 현재 구성종목의 현지통화 가격 단순수익률을 각 기간 시작일 시가총액으로 가중했습니다. BM은 해당 시장 전체 구성종목의 같은 방식으로 계산한 합성 수익률이며 공식 지수 수익률과 다를 수 있습니다. 배당과 환율 효과는 포함하지 않습니다. STOXX Europe 600은 시가총액 통화가 혼합돼 가중치와 합성 BM이 근사치입니다. 1D와 5D는 유효 거래일 기준이며, 1M·3M·6M은 해당 달력 기간 이전의 가장 가까운 유효 거래일 종가를 사용합니다. 데이터가 부족한 기간은 임의 보간하지 않고 ‘데이터 부족’으로 표시합니다.</div>
</section>"""


def render_combined_html(markets: list[dict[str, Any]], primary: str, descriptions: dict[tuple[str, str], str]) -> str:
    if not markets:
        raise ValueError("출력할 시장이 없습니다.")
    panels = "".join(render_market_panel(market, descriptions, idx == 0) for idx, market in enumerate(markets))
    buttons = "".join(
        f'<button class="market-tab{" active" if idx == 0 else ""}" data-market="{esc(market["key"])}" onclick="showMarket(\'{esc(market["key"])}\', this)">{esc(market["title"])}</button>'
        for idx, market in enumerate(markets)
    )
    as_of = markets[0]["asOf"]
    return f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>글로벌 섹터·산업 {primary} 보고서</title>
<style>
:root{{--navy:#17365d;--blue:#0070c0;--red:#c00000;--line:#d9e2f3;--pale:#f3f7fb;--text:#172033}}
*{{box-sizing:border-box}} body{{margin:0;font-family:Arial,'Malgun Gothic',sans-serif;color:var(--text);background:#fff}}
main{{max-width:1440px;margin:0 auto;padding:20px 24px 50px}} h1{{margin:0 0 5px;font-size:27px}} h2{{margin:26px 0 9px;padding-bottom:5px;border-bottom:2px solid var(--navy);font-size:19px}} h3{{margin:22px 0 8px;color:var(--navy);font-size:17px}} h4{{margin:13px 0 6px;font-size:14px}}
.meta{{font-size:12px;color:#667085;line-height:1.6}} .badge{{display:inline-block;padding:4px 10px;margin-right:5px;border-radius:999px;background:var(--navy);color:#fff;font-weight:700}}
.report-header{{margin-bottom:12px}} .report-header h1{{font-size:22px}} .market-tabs{{position:sticky;top:0;z-index:20;display:flex;gap:5px;padding:8px 0;background:rgba(255,255,255,.96);border-bottom:1px solid #cbd5e1;margin-bottom:16px}} .market-tab{{border:1px solid #aebdce;border-radius:5px;background:#fff;color:var(--navy);padding:6px 11px;font-size:12px;font-weight:700;cursor:pointer}} .market-tab.active{{background:var(--navy);color:#fff;border-color:var(--navy)}} .market-panel{{display:none}} .market-panel.active{{display:block}}
.table-wrap{{display:block;width:max-content;max-width:100%;overflow:auto;border:1px solid #b8c5d5;border-radius:5px}} table{{width:auto;min-width:700px;border-collapse:collapse;font-size:12px;white-space:nowrap}} th{{position:sticky;top:0;background:var(--navy);color:#fff;padding:5px 8px;text-align:left;line-height:1.25;z-index:1;border:1px solid #3c587b}} td{{padding:5px 8px;line-height:1.25;border:1px solid #ccd6e3}} tbody tr:nth-child(even){{background:var(--pale)}} tr:hover{{background:#fff6d8!important}} th.num,.num{{text-align:right;font-variant-numeric:tabular-nums}} .pos{{color:var(--blue);font-weight:700}} .neg{{color:var(--red);font-weight:700}} .zero{{color:#555}} .na{{color:#98a2b3;font-weight:400}} .benchmark{{background:#e2eaf4!important;border-top:2px solid var(--navy);font-weight:700}} .ticker{{color:#475467}}
.stock-name{{font-weight:700}} .stock-desc{{max-width:390px;margin-top:2px;color:#667085;font-size:10px;line-height:1.35;white-space:pre-line}}
.commentary-slot:empty{{display:none}} .commentary-block{{margin:8px 0;padding:9px 11px;border-left:3px solid var(--navy);background:#f8fafc;font-size:11px;line-height:1.55;white-space:normal}} .commentary-block strong{{display:block;margin-bottom:2px;color:var(--navy)}} .commentary-block .commentary-meta{{margin-top:3px;color:#667085;font-size:10px}} .stock-desc+.commentary-block{{max-width:390px;margin-top:5px}}
.cards{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px;margin:12px 0}} .card{{padding:13px 15px;border:1px solid var(--line);border-radius:8px;background:#fafcff}} .card span{{display:block;font-size:11px;color:#667085}} .card b{{display:block;margin:5px 0}} .card strong{{font-size:18px}}
.note{{margin-top:30px;padding:12px 14px;border-left:4px solid var(--navy);background:var(--pale);font-size:11px;line-height:1.65;color:#596579}} .sector-detail{{page-break-inside:avoid}}
@media(max-width:800px){{main{{padding:18px 12px}}.cards{{grid-template-columns:1fr}}}}
@media print{{main{{max-width:none;padding:10mm}}.table-wrap{{overflow:visible}}th{{position:static}}}}
</style></head><body><main data-report="sector" data-as-of="{esc(as_of)}" data-primary="{esc(primary)}">
<header class="report-header"><h1>글로벌 섹터·산업 로테이션 — {primary}</h1><div class="meta">기준일 {esc(as_of)} · 시장 탭을 눌러 이동</div></header>
{commentary_slot("global")}
<nav class="market-tabs" aria-label="시장 선택">{buttons}</nav>
{panels}
</main>
<script>
function showMarket(key, button) {{
  document.querySelectorAll('.market-panel').forEach(panel => panel.classList.toggle('active', panel.dataset.market === key));
  document.querySelectorAll('.market-tab').forEach(tab => tab.classList.toggle('active', tab === button));
  window.scrollTo({{top:0, behavior:'instant'}});
}}
</script></body></html>"""


def sanitize_for_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: sanitize_for_json(item) for key, item in value.items() if key != "indices"}
    if isinstance(value, list):
        return [sanitize_for_json(item) for item in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def main() -> None:
    args = parse_args()
    if not args.master.exists():
        raise FileNotFoundError(args.master)
    if not 1 <= args.top_n <= 10:
        raise ValueError("--top-n은 1~10 사이여야 합니다.")

    selected = list(MARKETS) if args.market == "ALL" else [args.market]
    workbook = openpyxl.load_workbook(args.master, read_only=True, data_only=True)
    try:
        missing = [MARKETS[key]["sheet"] for key in selected if MARKETS[key]["sheet"] not in workbook.sheetnames]
        if missing:
            raise KeyError(f"Master에 필요한 시트가 없습니다: {missing}")
        markets = [build_market(load_market(workbook, key), args.primary, args.top_n, args.as_of) for key in selected]
    finally:
        workbook.close()

    as_of_dates = {market["asOf"] for market in markets}
    if len(as_of_dates) != 1:
        raise ValueError(f"시장별 기준일이 일치하지 않습니다: {sorted(as_of_dates)}")
    storage_tag = markets[0]["asOf"].replace("-", "")
    output_dir = args.output_root / storage_tag / "Sector_Industry"
    output_dir.mkdir(parents=True, exist_ok=True)
    descriptions = load_descriptions(args.descriptions.resolve())

    html_path = output_dir / f"글로벌_섹터_산업_{args.primary}_보고서_{storage_tag}.html"
    json_path = output_dir / f"글로벌_섹터_산업_{args.primary}_데이터_{storage_tag}.json"
    html_path.write_text(render_combined_html(markets, args.primary, descriptions), encoding="utf-8")
    json_path.write_text(
        json.dumps(
            {"primary": args.primary, "asOf": markets[0]["asOf"], "markets": sanitize_for_json(markets)},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    outputs = [str(html_path.resolve())]

    print(json.dumps({"primary": args.primary, "asOf": markets[0]["asOf"], "outputs": outputs}, ensure_ascii=False))


if __name__ == "__main__":
    main()
