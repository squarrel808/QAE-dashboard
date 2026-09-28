from pathlib import Path
import sys, json, pickle
from datetime import date
from collections import Counter
import numpy as np
import openpyxl

ROOT=Path(r'C:\Users\infomax\Documents\python\BQL')
HERE=Path(__file__).resolve().parent
OUT=HERE.parent if HERE.name=='code' else HERE/'daily_20260907'
OUT.mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/'유로존_dashboard'))
sys.path.insert(0,str(ROOT/'Theme'/'output'))
import build_market_rotation_dashboard as b
import build_report_datasets as d
TARGETS=['SPX','SHSZ300','SXXP','NDX']
new=ROOT/'Rawfile'/'Daily_Input'/'Bquant_All_19_Raw_Refreshed12시 (5).xlsb'
raws={}
wb=openpyxl.load_workbook(ROOT/'Rawfile'/'BQuant_Master.xlsx',read_only=True,data_only=True)
for sheet in wb.sheetnames:
    if sheet.split(' ')[0] not in TARGETS: continue
    print('READ',sheet,flush=True)
    raw=b.parse_raw_rows([list(r) for r in wb[sheet].iter_rows(values_only=True)],'BQuant_Master.xlsx',sheet)
    if raw is not None:raws[raw.code]=raw
wb.close()
for raw in b.iter_raw_indices(new):
    if raw.code in TARGETS:raws[raw.code]=b.merge_raw_indices(raws[raw.code],raw)
audit=[]
for code,raw in raws.items():
    for j,dt in enumerate(raw.dates):
        if dt<date(2026,9,2):continue
        a=raw.metrics['price']; audit.append({'index':code,'date':str(dt),'valid':int(np.isfinite(a[:,j]).sum()),'changed':int(np.sum(np.isfinite(a[:,j])&np.isfinite(a[:,j-1])&~np.isclose(a[:,j],a[:,j-1])))})
    # Sep 5-7 prices repeat Sep 4 in this input; use completed, changed observations.
    keep=[j for j,dt in enumerate(raw.dates) if dt<=date(2026,9,4)]
    raw.dates=[raw.dates[j] for j in keep]
    raw.metrics={k:a[:,keep] for k,a in raw.metrics.items()}
with (OUT/'raws.pkl').open('wb') as f:pickle.dump(raws,f)
markets=[d.market_dataset(raws[c]) for c in TARGETS[:3]]
for m in markets:
    for key in ['1d','10d']:
        ranked=sorted(m['all'],key=lambda x:x['return'+key],reverse=True)
        m['top100_'+key]=ranked[:100];m['bottom100_'+key]=list(reversed(ranked[-100:]))
    # Use the previous session's market cap for a one-day return proxy.
    raw=raws[m['code']];ss=d.actual_sessions(raw.metrics['price']); e,p=ss[-1],ss[-2]
    r=d.simple_return(raw.metrics['price'],p,e);w=raw.metrics['mcap'][:,p]
    valid=np.isfinite(r)&np.isfinite(w)&(w>0)&~d.corporate_action_mask(raw.metrics['price'],raw.metrics['mcap'],ss[-11],p,e)
    m['marketReturn1dProxy']=float(np.average(r[valid],weights=w[valid]))
    m['sectors']=[]
    for s in sorted(set(raw.sectors)):
        mask=valid&np.array([x==s for x in raw.sectors]); rr=[x for x in m['all'] if x['sector']==s]
        if mask.any():m['sectors'].append({'sector':s,'n':len(rr),'return1d':float(np.average(r[mask],weights=w[mask])),'breadth':sum(x['return1d']>0 for x in rr)/len(rr) if rr else None})
ai=d.ai_dataset(raws['SPX'],raws['NDX'])
data={'report_date':'2026-09-07','cutoff_note':'신규 파일의 9월 5~7일 가격이 반복되어 9월 4일 마지막 유효 관측일 사용.','audit':audit,'markets':markets,'ai':ai}
(OUT/'data.json').write_text(json.dumps(data,ensure_ascii=False,allow_nan=False),encoding='utf-8')
for m in markets:
    print('\nMARKET',m['code'],m['previous'],m['end'],m['members'],'PROXY',m['marketReturn1dProxy'],m['marketReturn10dProxy'],flush=True)
    for label in ['best1d','worst1d','best10d','worst10d']:
        print(label,[(x['ticker'],x['name'],round(x['return1d']*100,2),round(x['return10d']*100,2)) for x in m[label]],flush=True)
print('AI',ai['meta'],flush=True)
print('STAGES',[(s['name'],round(s['return1d']*100,2),round(s['return1w']*100,2),round(s['return1m']*100,2),round(s['return3m']*100,2)) for s in ai['stages']],flush=True)
print('AUDIT',audit,flush=True)
