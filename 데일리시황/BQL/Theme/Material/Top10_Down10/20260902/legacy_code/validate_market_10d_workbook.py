import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
MANIFEST = HERE / "output" / "latest_outputs.json"
expected = {"Summary", "US_SPX", "China_CSI300", "Europe_STOXX600", "Methodology"}


def default_path():
    if not MANIFEST.exists():
        raise FileNotFoundError(f"Run build_market_10d_outputs.py first: {MANIFEST}")
    metadata = json.loads(MANIFEST.read_text(encoding="utf-8"))
    return Path(metadata["workbook"])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", type=Path)
    args = parser.parse_args()
    path = (args.workbook or default_path()).resolve()

    wb_formula = load_workbook(path, data_only=False, read_only=False)
    wb_values = load_workbook(path, data_only=True, read_only=True)
    assert set(wb_formula.sheetnames) == expected, wb_formula.sheetnames

    formula_count = 0
    errors = []
    for ws in wb_formula.worksheets:
        values_ws = wb_values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    cached = values_ws[cell.coordinate].value
                    if isinstance(cached, str) and cached.startswith("#"):
                        errors.append((ws.title, cell.coordinate, cached))

    for sheet in ["US_SPX", "China_CSI300", "Europe_STOXX600"]:
        ws = wb_formula[sheet]
        assert ws.max_row >= 37
        assert ws["I7"].data_type == "f"
        assert ws["I26"].data_type == "f"
        values = wb_values[sheet]
        for row in range(7, 27):
            expected_return = values[f"H{row}"].value / values[f"G{row}"].value - 1
            assert abs(values[f"I{row}"].value - expected_return) < 1e-10

    assert formula_count == 60, formula_count
    assert not errors, errors
    print(f"sheets={wb_formula.sheetnames}")
    print(f"formulas={formula_count}")
    print("formula_errors=0")
    print(path)


if __name__ == "__main__":
    main()
