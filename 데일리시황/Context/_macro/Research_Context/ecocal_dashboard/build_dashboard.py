"""Read-only calendar/research ingestion; reproducible local dashboard output."""
from __future__ import annotations
import argparse,collections,hashlib,importlib.util,json,math,re,sqlite3,sys
from datetime import datetime,date
from pathlib import Path
import openpyxl

HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE))
import indicator_history
DEFAULT_ROOT=Path(r'C:\Users\infomax\Documents\python\QAE\데일리시황\Context\_macro\Research_Context')
DEFAULT_RAW=Path(r'C:\Users\infomax\Documents\python\QAE\____Rawdata___')
HOUSES=['GS','JPM','Citi','BofA','HSBC']
COUNTRIES={'United States':('US','미국'),'Japan':('JP','일본'),'Germany':('DE','독일'),'France':('FR','프랑스'),'Italy':('IT','이탈리아'),'United Kingdom':('GB','영국'),'Canada':('CA','캐나다'),'Russia':('RU','러시아'),'Euro Area':('EA','유로존')}
COUNTRIES.update({'South Korea':('KR','한국'),'China':('CN','중국'),'Australia':('AU','호주')})
COUNTRY_ALIASES={code:name for name,(code,_) in COUNTRIES.items()}|{'JN':'Japan','UK':'United Kingdom','GE':'Germany','EC':'Euro Area','Eurozone Aggregate':'Euro Area'}
THEMES={'growth':'경기','inflation':'물가','policy':'통화정책'}
FIELDS=['country','date','time','event','period','survey','actual','prior','revised','scale']
LABELS=['국가','발표일','발표시각','지표','대상 기간','컨센서스','실제값','종전값','수정 종전값','단위']
EXPORT_COLUMNS=[{'letter':chr(65+i),'key':key,'label':label} for i,(key,label) in enumerate(zip(
    ['datetime','country_code','index','blank','event','period','survey','actual','prior','revised','relevance','ticker'],
    ['발표일시','국가 코드','순번','빈 열','지표','대상 기간','컨센서스','실제값','종전값','수정 종전값','중요도','티커']))]

# Indicator-first search is deliberately separate from the broad living-IB views.
# These aliases find pages for review; a hit is never treated as a verified IB comment.
GEO_ALIASES={
    'US':['United States','U.S.','US'], 'JP':['Japan','Japanese'],
    'DE':['Germany','German'], 'FR':['France','French'], 'IT':['Italy','Italian'],
    'GB':['United Kingdom','UK','British'], 'CA':['Canada','Canadian'],
    'RU':['Russia','Russian'], 'EA':['Euro area','Eurozone','ECB'],
    'KR':['South Korea','Korea','Korean'], 'CN':['China','Chinese'], 'AU':['Australia','Australian']}

def norm(s):return re.sub(r'\s+',' ',re.sub(r'(?m)^\s*(?:n|o|§|◆|•)\s*$',' ',s.replace('\u00ad',''))).strip()
def numeric(x):return float(x) if type(x) in (int,float) and math.isfinite(x) else None
def fmt(x):return '미수록' if x is None else f'{x:,.5f}'.rstrip('0').rstrip('.')
def identity(e):return tuple(str(e[k]) for k in FIELDS[:5])
def reference(period,release):
    if not period or not period.strip():return None
    p=period.split()[0];year=int(release[:4])
    months={v:i for i,v in enumerate(['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'],1)}
    if p in months:
        m=months[p];year-=int(m>int(release[5:7]));return f'{year}-{m:02d}'
    if re.fullmatch(r'[1-4]Q',p):
        year-=int(int(p[0])>(int(release[5:7])-1)//3+1)
        return f'{year}-Q{p[0]}'
    return period

def classify(name,country,scale):
    n=name.lower();family='activity';theme='growth';note=''
    if any(t in n for t in ['rate decision','bank rate','key rate','target rate','deposit facility rate','main refinancing rate','marginal lending facility']):theme='policy';family='policy_rate'
    elif any(t in n for t in ['money stock','tic flows','foreign bonds','forex reserve']):
        theme='policy';family='liquidity';note='통화정책의 보조 지표: 자금흐름·유동성이며 중앙은행 결정 자체는 아님.'
    elif any(t in n for t in ['cpi','hicp','ppi','deflator','retail price index','rpi ','import price','industrial product price','inflation expectations']):theme='inflation';family='prices'
    elif any(t in n for t in ['earnings','wage']):theme='growth';family='wages';note='주분류는 경기·소득이며 물가·통화정책에도 연결.'
    elif any(t in n for t in ['payroll','unemployment','employment','jobless','continuing claims','claimant']):family='labor'
    elif 'gdp' in n:family='gdp'
    elif any(t in n for t in ['pmi','ism services']):family='pmi'
    elif any(t in n for t in ['production','orders','capacity utilization']):family='production'
    elif any(t in n for t in ['retail sales']):family='consumption'
    elif any(t in n for t in ['home sales','housing','building permits','mortgage','house price']):family='housing'
    elif any(t in n for t in ['trade balance','current account']):family='trade'
    elif 'inventories' in n:family='inventories'
    elif 'budget' in n:family='fiscal';note='재정·국채 공급과 연결하는 경기 보조 지표.'
    unit='index' if scale is None else {'k':'thousand','m':'million','b':'billion','t':'trillion'}.get(scale,scale)
    if scale=='%':
        unit='pct_yoy' if 'yoy' in n else 'pct_mom' if 'mom' in n else 'pct_qoq_annualized' if 'annualized' in n and 'qoq' in n else 'pct_qoq' if 'qoq' in n else 'pct_wow' if 'wow' in n else 'pct_level'
    suffix=''
    if family=='prices':
        if 'cpi' in n:suffix='CORE_CPI' if 'core' in n or 'ex food' in n else 'HEADLINE_CPI'
        elif 'ppi' in n:suffix='CORE_PPI' if 'ex food' in n else 'PPI'
        elif 'deflator' in n:suffix='GDP_DEFLATOR'
    if family=='gdp':suffix='REAL_GDP'
    if n=='change in nonfarm payrolls':suffix='PAYROLLS'
    if n=='unemployment rate':suffix='UNEMPLOYMENT'
    watches={
        'prices':['CPI','CORE_CPI','HICP','CORE_HICP','PPI','CORE_PCE'], 'wages':['WAGES','AHE','CONSUMPTION'],
        'labor':['EMPLOYMENT','PAYROLLS','UNEMPLOYMENT','PARTICIPATION'], 'gdp':['GDP','REAL_GDP','CONSUMPTION'],
        'pmi':['PMI','GDP'],'production':['GDP','CAPEX'],'consumption':['RETAIL_SALES','CONSUMPTION','GDP'],
        'housing':['HOUSE_PRICES','GDP'],'trade':['GDP'],'inventories':['GDP'],'fiscal':['FISCAL','GDP'],
        'liquidity':['FISCAL'],'policy_rate':[],'activity':['GDP']}
    regional='EA' if country in {'DE','FR','IT'} else country
    watch=[regional+'_'+x for x in watches[family]]
    decision={'US':'FED_DECISION','GB':'BOE_DECISION','JP':'BOJ_DECISION','CA':'BOC_DECISION','EA':'ECB_DECISION'}.get(regional)
    if decision and theme=='policy':watch.append(decision)
    return dict(theme=theme,family=family,classification_note=note,unit=unit,indicator=country+'_'+suffix if suffix else None,watch=watch)


def _fts_group(terms):
    clean=[]
    for term in terms:
        term=' '.join(str(term).strip().split())
        if term and term.lower() not in {x.lower() for x in clean}:clean.append(term)
    return '('+' OR '.join('"'+x.replace('"','""')+'"' for x in clean)+')'


def indicator_search_spec(e):
    """Return an indicator-specific FTS query and the aliases used to audit it."""
    n=e['event'].lower();family=e['family'];base=[];detail=[]
    if family=='pmi':
        base=['PMI','purchasing managers','ISM']
        if 'service' in n or 'nonmanufactur' in n:detail=['services','nonmanufacturing','business activity']
        elif 'manufactur' in n:detail=['manufacturing','factory']
        elif 'composite' in n:detail=['composite','all-industry','all industry']
    elif family=='labor':
        if 'payroll' in n:base=['nonfarm payrolls','payrolls','employment report']
        elif 'unemployment' in n or 'claimant' in n:base=['unemployment rate','claimant count','labor market']
        elif 'jobless' in n or 'claims' in n:base=['jobless claims','unemployment claims','initial claims','continuing claims']
        elif 'jolt' in n or 'job opening' in n:base=['JOLTS','job openings']
        elif 'adp' in n:base=['ADP employment','ADP payrolls']
        else:base=['employment','labor market','labour market']
    elif family=='wages':base=['wages','earnings','average hourly earnings','labor income','labour income']
    elif family=='prices':
        if 'inflation expectations' in n:base=['inflation expectations','survey of consumer expectations']
        elif 'ppi' in n or 'producer' in n:base=['PPI','producer prices','producer price index']
        elif 'pce' in n:base=['PCE inflation','PCE price index','PCE deflator']
        elif 'rpi' in n or 'retail price index' in n:base=['RPI','retail price index']
        elif 'import price' in n:base=['import prices','import price index']
        elif 'deflator' in n:base=['GDP deflator','implicit price deflator']
        else:base=['CPI','consumer prices','consumer price index','HICP']
        if 'core' in n or 'ex food' in n:detail=['core','excluding food','ex food']
        elif 'headline' in n:detail=['headline']
    elif family=='gdp':base=['GDP','gross domestic product'];detail=['growth','output']
    elif family=='production':
        if 'order' in n:base=['factory orders','durable goods orders','orders']
        elif 'capacity' in n:base=['capacity utilization','capacity utilisation']
        else:base=['industrial production','manufacturing production','factory output']
    elif family=='consumption':base=['retail sales','consumer spending','consumption']
    elif family=='housing':
        if 'mortgage' in n:base=['mortgage applications','MBA mortgage','mortgage demand']
        elif 'permit' in n:base=['building permits','housing permits']
        elif 'start' in n:base=['housing starts','homebuilding']
        elif 'home sales' in n:base=['home sales','house sales']
        else:base=['house prices','home prices','housing market','RICS']
    elif family=='trade':base=['trade balance','exports','imports','external trade']
    elif family=='inventories':base=['inventories','inventory','stockbuilding']
    elif family=='fiscal':base=['budget balance','fiscal balance','government deficit']
    elif family=='liquidity':base=['money supply','money stock','capital flows','reserves']
    elif family=='policy_rate':base=['rate decision','policy rate','central bank'];detail=['meeting','decision']
    else:
        tokens=[x for x in re.findall(r'[A-Za-z][A-Za-z-]{2,}',e['event']) if x.lower() not in {'index','rate','change','final','preliminary','seasonally','adjusted'}]
        base=[' '.join(tokens[:5])] if tokens else [e['event']]
    query=_fts_group(base)
    if detail:query+=' AND '+_fts_group(detail)
    return {'query':query,'aliases':base,'qualifiers':detail}


def _candidate_excerpt(text,terms,limit=1000):
    text=norm(text);lower=text.lower();hits=[lower.find(t.lower()) for t in terms if lower.find(t.lower())>=0]
    pos=min(hits) if hits else 0;start=max(0,pos-220);end=min(len(text),pos+limit)
    excerpt=text[start:end]
    for marker in ['For the exclusive use of','Unauthorized redistribution']:
        if marker in excerpt:excerpt=excerpt.split(marker,1)[0]
    excerpt=re.sub(r'\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b','[email omitted]',excerpt)
    return ('…' if start else '')+excerpt.strip()+('…' if end<len(text) else '')


def indicator_candidates(src,e,house,asof,limit=3):
    """Find geographically scoped pages mentioning this indicator.

    Results are unreviewed search candidates. They must be read in full before an
    IB summary is written. Broad growth/inflation views are not returned here.
    """
    spec=indicator_search_spec(e);region='EA' if e['country_code'] in {'DE','FR','IT'} else e['country_code']
    scopes=[('country',e['country_code'],0)]
    if region!=e['country_code']:scopes.append(('region',region,1))
    rows=[]
    sql='''SELECT d.doc_id,d.house,d.title,d.published,c.page,c.text,bm25(chunks_fts) AS rank
           FROM chunks_fts f JOIN chunks c ON c.chunk_id=f.rowid JOIN documents d ON d.doc_id=c.doc_id
           WHERE d.status='indexed' AND d.date_review=0 AND d.house=? AND d.published<=?
             AND chunks_fts MATCH ?
             AND EXISTS(SELECT 1 FROM tags t WHERE t.doc_id=d.doc_id AND t.kind='country' AND t.value=?)
           ORDER BY rank,d.published DESC LIMIT 12'''
    for scope,code,priority in scopes:
        for row in src.execute(sql,(house,asof,spec['query'],code)):
            rows.append((priority,scope,dict(row)))
    # Some global reports contain a country section but only carry a global tag.
    # Keep this fallback geographically explicit by requiring the country name in FTS.
    if len(rows)<limit:
        geo=GEO_ALIASES.get(e['country_code'],[e['country']])
        geo_query=spec['query']+' AND '+_fts_group(geo)
        fallback=src.execute(sql.replace("AND EXISTS(SELECT 1 FROM tags t WHERE t.doc_id=d.doc_id AND t.kind='country' AND t.value=?)",''),(house,asof,geo_query)).fetchall()
        rows.extend((2,'country_text',dict(r)) for r in fallback)
        if region!=e['country_code'] and len(rows)<limit:
            regional_query=spec['query']+' AND '+_fts_group(GEO_ALIASES['EA'])
            fallback=src.execute(sql.replace("AND EXISTS(SELECT 1 FROM tags t WHERE t.doc_id=d.doc_id AND t.kind='country' AND t.value=?)",''),(house,asof,regional_query)).fetchall()
            rows.extend((3,'region_text',dict(r)) for r in fallback)
    found=[];seen=set();release=date.fromisoformat(e['date'])
    for priority,scope,row in rows:
        key=(row['doc_id'],row['page'])
        if key in seen:continue
        seen.add(key);published=date.fromisoformat(row['published']);delta=(published-release).days
        timing='same_day_timing_unverified' if delta==0 else 'after_release_candidate' if delta>0 else 'pre_release_candidate'
        paths=[r[0] for r in src.execute('select path from files where doc_id=? order by path',(row['doc_id'],))]
        available=next((p for p in paths if Path(p).is_file()),None)
        if not available:continue
        title_hit=any(x.lower() in row['title'].lower() for x in spec['aliases']+spec['qualifiers'])
        found.append({k:row[k] for k in ('doc_id','house','title','published','page')}|{
            'match_scope':scope,'timing':timing,'days_from_release':delta,'title_hit':title_hit,
            'search_query':spec['query'],'search_aliases':spec['aliases'],'search_qualifiers':spec['qualifiers'],
            'excerpt':_candidate_excerpt(row['text'],spec['aliases']+spec['qualifiers']),
            'path':available,'url':Path(available).as_uri()+'#page='+str(row['page']),
            '_sort':(priority,0 if 0<=delta<=14 else 1 if -21<=delta<0 else 2,not title_hit,abs(delta),row['rank'])})
    found.sort(key=lambda x:x['_sort'])
    for item in found:item.pop('_sort',None)
    return found[:limit]

def parse_export_number(value,number_format='General'):
    """Return displayed magnitude, explicit unit and currency; reject unknown text."""
    if value is None or isinstance(value,str) and value.strip() in {'','--'}:return None,None,None
    display_format=re.sub(r'"[^"]*"|\\.', '', number_format or '')
    if type(value) in (int,float):
        if not math.isfinite(value):raise ValueError(f'유한 수치가 아님: {value!r}')
        return (round(float(value)*100,12),'%',None) if '%' in display_format else (float(value),None,None)
    if not isinstance(value,str):raise ValueError(f'지원하지 않는 수치 형식: {value!r}')
    match=re.fullmatch(r'([+-]?)([$£¥€]?)([+-]?)(\d+(?:,\d{3})*(?:\.\d+)?|\.\d+)([%kmbt]?)',value.strip())
    if not match or match[1] and match[3]:raise ValueError(f'지원하지 않는 수치/단위: {value!r}')
    sign=match[1] or match[3];number=float((sign or '')+match[4].replace(',',''))
    if not math.isfinite(number):raise ValueError(f'유한 수치가 아님: {value!r}')
    if match[2] and match[5]=='%':raise ValueError(f'통화/퍼센트 단위 충돌: {value!r}')
    return number,match[5] or None,match[2] or None

def parse_export_datetime(value):
    if isinstance(value,datetime):return value.date().isoformat(),value.strftime('%H:%M:%S' if value.second else '%H:%M')
    if isinstance(value,date):return value.isoformat(),''
    if isinstance(value,str):
        for pattern in ('%m/%d/%Y','%Y-%m-%d'):
            try:return datetime.strptime(value.strip(),pattern).date().isoformat(),''
            except ValueError:pass
    raise ValueError(f'지원하지 않는 발표일시: {value!r}')

def export_event(cells):
    values=[c.value for c in cells];country=COUNTRY_ALIASES.get(values[1],values[1])
    if not values[1] and not values[4]:return None
    if country not in COUNTRIES or not isinstance(values[4],str) or not values[4].strip():raise ValueError(f'국가/지표 누락 또는 미지원: {values[1:5]}')
    released,clock=parse_export_datetime(values[0]);period=values[5]
    period=f'{period:%b} {period.day}' if isinstance(period,(date,datetime)) else '' if period is None else str(period)
    parsed=[parse_export_number(c.value,c.number_format) for c in cells[6:10]]
    units={unit for value,unit,currency in parsed if value is not None};currencies={currency for value,unit,currency in parsed if value is not None}
    if len(units)>1 or len(currencies)>1:raise ValueError(f'한 지표의 수치 단위/통화 불일치: {values[4]} {parsed}')
    return dict(country=country,date=released,time=clock,event=values[4].strip(),period=period,
                **dict(zip(['survey','actual','prior','revised'],[p[0] for p in parsed])),
                scale=next(iter(units),None),currency=next(iter(currencies),None),
                raw={chr(65+i):v.isoformat() if isinstance(v,(date,datetime)) else v for i,v in enumerate(values)},
                source_layout='bloomberg_export',source_column_range='A:L',source_columns=EXPORT_COLUMNS)

def load_calendar(path):
    workbook=openpyxl.load_workbook(path,read_only=True,data_only=True);events=[];skipped=[];seen=set()
    for sheet in workbook:
        if sheet.max_column<11:continue
        header=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
        exported=len(header)>=12 and header[0]=='Date Time' and header[1]=='Country Code' and header[4]=='Event'
        for row,cells in enumerate(sheet.iter_rows(min_row=2 if exported else 3,max_col=12 if exported else 11),2 if exported else 3):
            if exported:
                try:e=export_event(cells)
                except ValueError as exc:raise ValueError(f'{sheet.title}:{row}: {exc}') from exc
                if e is None:continue
            else:
                raw=[c.value for c in cells[1:11]]
                if not any(v is not None for v in raw):continue
                raw+=[None]*(10-len(raw));e=dict(zip(FIELDS,raw))
                e['raw']={chr(66+i):v.isoformat() if isinstance(v,(date,datetime)) else v for i,v in enumerate(raw)}
            e['country']=COUNTRY_ALIASES.get(e['country'],e['country'])
            if e['country'] not in COUNTRIES:raise ValueError(f'국가 매핑 필요: {e["country"]}, {sheet.title}:{row}')
            if isinstance(e['date'],(date,datetime)):e['date']=e['date'].strftime('%Y-%m-%d')
            date.fromisoformat(str(e['date']));e['date']=str(e['date']);e['time']=str(e['time']);e['period']=str(e['period'])
            code,label=COUNTRIES[e['country']];e.update(country_code=code,country_label=label)
            ident=identity(e)
            if ident in seen:raise ValueError(f'중복 지표 식별자: {ident}')
            seen.add(ident);e['id']=hashlib.sha256('|'.join(ident).encode()).hexdigest()[:20]
            for k in ['survey','actual','prior','revised']:e[k]=numeric(e[k])
            e.update(source_sheet=sheet.title,source_row=row,release_month=e['date'][:7],reference_month=reference(e['period'],e['date']))
            e.update(classify(e['event'],code,e['scale']))
            e['surprise']=None if e['actual'] is None or e['survey'] is None else round(e['actual']-e['survey'],8)
            e['revision']=None if e['revised'] is None or e['prior'] is None else round(e['revised']-e['prior'],8)
            e['baseline']=e['revised'] if e['revised'] is not None else e['prior']
            e['actual_change']=None if e['actual'] is None or e['baseline'] is None else round(e['actual']-e['baseline'],8)
            e['status']='reported' if e['actual'] is not None else 'missing'
            events.append(e)
    workbook.close();return sorted(events,key=lambda e:(e['date'],e['time'],e['source_row']))

def evidence(src,e):
    doc=src.execute('select * from documents where doc_id=?',(e['doc_id'],)).fetchone()
    p=src.execute('select text from pages where doc_id=? and page=?',(e['doc_id'],e['page'])).fetchone()
    if not doc or not p or doc['status']!='indexed' or doc['date_review'] or norm(e['quote']) not in norm(p['text']):raise ValueError(f'인용 검증 실패: {e}')
    paths=[r[0] for r in src.execute('select path from files where doc_id=? order by path',(e['doc_id'],))]
    available=next((p for p in paths if Path(p).is_file()),None)
    if not available:raise ValueError(f'원본 PDF 없음: {doc["title"]}')
    return {**e,'title':doc['title'],'published':doc['published'],'house':doc['house'],'path':available,'url':Path(available).as_uri()+'#page='+str(e['page'])}

def factual(e):
    if e['actual'] is None:
        return f'파일 실제값 미수록. 컨센서스 {fmt(e["survey"])} / '+('수정 종전 ' if e['revised'] is not None else '종전 ')+fmt(e['baseline'])+'.'
    s=f'실제 {fmt(e["actual"])}'
    if e['surprise'] is not None:s+=f', 예상 대비 {e["surprise"]:+g}'+('%p' if e['scale']=='%' else ' '+(e['scale'] or '포인트'))
    if e['baseline'] is not None:
        revision_release=bool(re.search(r'\b[FP]\b',e['period'])) and (e['family']=='gdp' or e['family']=='pmi' or 'CPI' in e['event'] or e['family']=='production' or e['family']=='inventories')
        s+=f'. '+('원본 Prior(직전 추정치일 수 있음) ' if revision_release else '수정 종전 ' if e['revised'] is not None else '종전 ')+fmt(e['baseline'])+f' 대비 {e["actual_change"]:+g}'
    if e['revision'] is not None:s+=f'. 과거값 수정 {e["revision"]:+g}'
    return s+'.'

def interpretation(e):
    future=e['actual'] is None;family=e['family'];n=e['event'].lower()
    if family=='prices':
        text='헤드라인·근원, 월간 속도·전년비, 에너지 기저효과를 나눠 물가의 지속성을 확인한다.'
        if 'ppi' in n or 'industrial product price' in n:text='생산자 단계의 원가 신호다. 소비자물가 전가와 마진 흡수를 별도로 봐야 하며 CPI와 같은 지표로 취급하지 않는다.'
        if 'deflator' in n:text='GDP 디플레이터는 국내 생산의 가격지표다. 소비자물가와 대상·가중치가 달라 CPI 전망에 그대로 대입하지 않는다.'
    elif family=='gdp':text='연율과 비연율을 구별하고, 소비·설비투자·정부·재고·순수출 기여를 확인해야 성장의 질을 판정할 수 있다. 확정치의 Prior는 같은 분기 속보치일 수 있다.'
    elif family=='labor':
        text='고용의 폭·근로시간·노동공급·임금을 함께 확인한다. 한 번의 고용 상회만으로 물가나 금리 경로를 확정하지 않는다.'
        if 'unemployment' in n or 'claimant count rate' in n:text='실업률 변화는 고용 수요와 경제활동참가율 모두의 영향을 받는다. 비반올림 수치와 참가율을 함께 확인해야 한다.'
        if 'claims' in n:text='청구 증가는 노동시장 약화 방향이지만 주간 변동성이 커 4주 평균·수정치와 재취업 흐름을 확인한다.'
    elif family=='pmi':
        text='50 기준 확장·위축과 신규주문·고용·가격의 구성을 함께 확인한다. 확정치와 속보치 차이를 전월 변화로 해석하지 않는다.'
        if e['actual'] is not None:text=f'지수 {fmt(e["actual"])}는 50 '+('위의 확장' if e['actual']>50 else '아래의 위축' if e['actual']<50 else '경계')+' 영역이다. '+text
    elif family=='wages':text='명목임금 상승은 소득을 지지하지만 실질 구매력 개선 여부는 물가와 함께 봐야 한다. 보너스와 기본급, 생산성 대비 임금비용을 구분한다.'
    elif family=='production':text='한 달 생산·수주는 자동차 조업·대형 주문·기저효과의 영향을 받을 수 있다. 업종별 확산과 3개월 추세로 지속성을 점검한다.'
    elif family=='consumption':text='소비 수요를 확인하되 명목 매출의 가격효과, 자동차 등 변동성이 큰 품목과 실질 소비를 구별한다.'
    elif family=='housing':text='금융여건에 민감한 주택 수요·공급 신호다. 거래·신규착공·허가·가격은 서로 다른 단계여서 동일한 경기 신호로 합산하지 않는다.'
    elif family=='trade':text='수출과 수입의 변화 및 가격·물량을 분해해야 한다. 적자 축소가 수입 위축 때문이면 내수 개선을 뜻하지 않으며 GDP 기여도와 경상수지도 구별한다.'
    elif family=='inventories':text='재고 증가가 최종수요를 앞선 비자발적 축적인지 확인한다. GDP 기여에는 재고 증감의 변화와 실질화가 필요하다.'
    elif family=='fiscal':text='월간 재정수지는 납세 일정·계절성 영향이 크다. 누적 적자와 국채 발행, 이자비용을 함께 봐야 금융여건 영향이 드러난다.'
    elif family=='liquidity':text='유동성·국제 자금흐름의 보조 신호다. 평가효과·거래 방향·자금 구성 없이 금리 인상·인하 의도로 해석하지 않는다.'
    elif family=='policy_rate':text='정책금리 수준을 직전 수준과 비교하고, 결정문·전망·표결·후속 가이던스로 다음 경로를 확인한다. 시장 컨센서스와 특정 IB 전망은 별개다.'
    else:text='심리·선행지표 신호가 실제 생산·소비·고용으로 이어지는지 후속 실적을 확인한다.'
    return ('발표 전 확인할 조건: ' if future else '이번 수치의 해석: ')+text

def select_views(views,e,house,asof):
    region='EA' if e['country_code'] in {'DE','IT','FR'} else e['country_code']
    candidates=[v for v in views if v['house']==house and v['country']==region and v['effective_date']<=asof]
    groups=collections.defaultdict(list)
    for v in candidates:groups[v['track']].append(v)
    latest=[]
    for track,history in groups.items():
        day=max(v['effective_date'] for v in history);tips=[v for v in history if v['effective_date']==day];superseded={v['prior_view_id'] for v in tips}
        latest.extend(v for v in tips if v['view_id'] not in superseded)
    def score(v):
        match=len(set(v['watch'])&set(e['watch']))
        growth_tracks={'labor':{'labor'},'wages':{'labor','productivity'},'housing':{'housing'},'consumption':{'consumption','growth'},'gdp':{'growth','consumption'},'production':{'growth'},'pmi':{'growth'},'trade':{'growth'},'inventories':{'growth'},'fiscal':{'rates'},'activity':{'growth'}}
        same=(e['family']=='policy_rate' and v['track'].startswith('policy.')) or (e['theme']=='inflation' and v['track'].startswith('inflation.')) or (e['theme']=='growth' and v['track'].split('.')[0] in growth_tracks.get(e['family'],set()))
        return (match>0,same,match,v['effective_date'])
    # No geographic-only fallback: a CPI view cannot become a trade-balance view.
    relevant=[v for v in latest if score(v)[0] or score(v)[1]]
    selected=sorted(relevant,key=score,reverse=True)[:2]
    if selected:
        background=sorted((v for v in latest if v['track'] in {'policy.near_term','policy.reaction'}),key=lambda v:(v['effective_date'],v['track']=='policy.near_term'),reverse=True)
        if background and all(v['view_id']!=background[0]['view_id'] for v in selected):selected.append({**background[0],'context_only':True})
    return selected,region


def stance_reading(e,house,bases,region):
    if not bases:return '관련성이 있는 저장 IB 견해가 부족해 하우스별 방향 판정은 보류한다.'
    family=e['family'];theme=e['theme'];future=e['actual'] is None
    macro={
      'US':{
       'growth':{'GS':'고용은 노동공급·계절성, GDP는 수출입 구성과 추적치를 함께 본다. 강한 지표 하나를 연준 인상으로 연결하기보다 근원물가 약 0.2%라는 동결 전제가 유지되는지 확인하는 해석이다.','JPM':'견조한 소비·노동시장이라는 전제에 비춰 경기 하방 위험을 점검한다. 성장 신호가 강해도 CPI가 결정적이며, 물가 판단이 팽팽할 때만 고용 강세가 매파 쪽에 힘을 더하는 구조다.','Citi':'고용 부진을 근거로 한 조기 인하 전제는 이미 약해졌다. 새로운 약세가 지속되면 인하 지연 전망을 재점검할 이유가 생기지만, 한 번의 약세로 인하 재개를 확정하지 않는다.','BofA':'AI 투자·소비와 견조한 성장이라는 75bp 인상 전망의 수요 전제를 시험한다. 광범위하고 지속적인 경기 약화는 이 전제를 약화시키며, 강한 실적은 이를 지지하지만 인상에는 물가 지속성 확인도 필요하다.','HSBC':'성장 회복력과 통제되는 근원물가의 조합에 비춰 본다. 성장 상회만으로 동결 전망을 버리기보다 임금·근원물가의 재가속이 동반되는지를 확인한다.'},
       'inflation':{'GS':'근원 CPI·PCE 약 0.2%가 동결의 주요 전제다. 같은 단위의 근원 지표가 지속적으로 웃돌면 이 전제가 약해지고, 부합·하회하면 동결 논거가 유지되는 방향이다.','JPM':'고용보다 이번 CPI를 결정적 정책 입력으로 본다. 물가 신호가 애매하면 강한 고용이 매파 판단을 보강하지만 PPI·헤드라인만으로 근원 CPI 경로를 확정할 수 없다.','Citi':'근원 CPI 0.184%·PCE 0.19%와 PCE 하향 수정이라는 디스인플레이션 전제를 시험한다. 지속적인 근원 상회는 동결 논거에 부담이지만 헤드라인 에너지 상승은 별도로 분리한다.','BofA':'근원 CPI 0.22%·PCE 0.24% 정도도 9월 인상에 충분하다고 보는 관점이다. 같은 0.2%대 숫자라도 GS·Citi와 정책 판단이 다르며, 예상보다 뚜렷하고 지속적인 둔화가 있어야 인상 전제가 약해진다.','HSBC':'에너지·식품의 헤드라인 위험과 통제되는 근원물가를 구별한다. 근원의 재가속 여부가 2026~2027년 동결 전제 및 인상 위험을 가르는 핵심이다.'},
       'policy':{'GS':'근원물가가 약 0.2%로 나오면 9월 동결을 뒷받침한다는 기존 조건을 결정문과 대조한다.','JPM':'저장된 견해는 CPI에 조건부인 반응함수다. 확정적 동결·인상 숫자가 확보되지 않아 금리 경로를 대신 만들어 넣지 않는다.','Citi':'9월 동결과 2027년 6월 인하 재개라는 기존 경로를 실제 결정·가이던스와 대조한다.','BofA':'9월부터 연내 총 75bp 인상이라는 전망을 대조한다. 파일 컨센서스의 동결과 IB 전망을 분리해 봐야 한다.','HSBC':'2026~2027년 동결이 기본이고 9월 인상 위험은 열려 있다는 관점이다. 기본 전망과 위험 시나리오를 따로 대조한다.'}},
      'JP':{
       'growth':{'GS':'민간소비·설비투자가 약한 GDP 구성과 다음 9월 인상 전망을 함께 본다. 총성장률 상향보다 민간 내수의 회복 여부가 중요하며 명목급여의 상회는 실질소득·기본급 지속성까지 확인해야 한다.','JPM':'정책에 따른 내구재 변동을 평탄화하면 소비 기조가 견조하다는 전제를 시험한다. 소득·서비스 소비 개선은 이를 지지하지만 GDP 헤드라인만으로 소비 기조를 판정하지 않는다.','Citi':'엔화·금융여건을 포함한 9월 인상 전제와 연결한다. 명목임금이나 한 차례 성장 하회만으로 정상화 경로가 뒤집혔다고 판단하지 않는다.','BofA':'분기당 한 차례에 가까운 인상 경로는 물가·임금의 지속성을 필요로 한다. 임금 상회는 소득과 비용 양쪽 신호이며 내수 약화가 지속되면 인상 속도의 부담이 된다.','HSBC':'저장된 일본 견해는 9월 인상이라는 글로벌 보고서의 짧은 요약이다. 개별 성장지표에 대한 반응함수와 이후 인상 경로는 근거가 부족해 방향 판정을 제한한다.'},
       'inflation':{'GS':'엔 약세와 완화적 금융여건을 배경으로 9월 인상을 앞당긴 견해다. 생산자물가 상회가 소비자가격·임금으로 이어지는지 확인해야 하며 PPI만으로 추가 인상 횟수를 늘리지 않는다.','JPM':'기존 9월·12월 인상 경로에 비춰 가격 압력의 전가와 지속성을 점검한다. GDP 디플레이터와 CPI를 같은 수치로 대입하지 않는다.','Citi':'9월 정상화 전망을 검토하되 PPI와 소비자 근원물가 사이 전가를 확인한다. 물가 한 항목으로 이후 횟수를 변경하지 않는다.','BofA':'식품 재가속과 신선식품·에너지 제외 물가 상승 전망을 갖고 있다. 이번 지표가 식품·임금·가격전가 중 어디를 확인하는지 따지며 PPI 상승을 곧바로 해당 근원물가 상승으로 간주하지 않는다.','HSBC':'9월 인상 요약 견해는 있지만 일본 고유의 물가 반응함수는 부족하다. 지표가 인상 전제를 강화했는지 확정하기에는 추가 원문이 필요하다.'}},
      'EA':{
       'growth':{'GS':'유로존 정책 견해는 9월 마지막 인상을 기본으로 두고 12월 위험을 열어둔다. 국가별 경기 약화의 지속성과 광범위함을 확인해야 추가 인상 위험을 재평가할 수 있다.','JPM':'서비스·소비의 회복력이 12월 추가 인상을 감당할 수 있다는 관점이다. 개별 제조업 부진이 서비스·소비로 확산되는지 확인해야 이 전제가 약해졌다고 볼 수 있다.','Citi':'확보한 유로존 견해는 다른 국가 보고서의 정책 캘린더 요약에 한정된다. 국가별 경기지표에서 금리 반응까지 추정할 직접 근거는 부족하다.','BofA':'상반기 이월효과와 달리 이후 분기 경로는 약해진다는 관점이다. 생산·수요 약화는 이 전망과 부합하는 방향이고, 지속적인 내수 상회는 이를 재점검하게 한다.','HSBC':'독일 재정·AI 투자·심리 회복의 지속력과 하반기 완만한 둔화를 함께 본다. 지표의 업종 구성과 외수·실질소득 부담을 분리해 회복의 폭을 확인한다.'},
       'inflation':{'GS':'근원물가 하회로 전망을 낮췄던 전제를 검토한다. 특정 국가 헤드라인 확정치만으로 유로존 전체 근원 경로를 상향하지 않는다.','JPM':'9월에 이어 12월 추가 인상이라는 견해다. 국가별 가격 신호가 유로존 서비스·임금·2차 효과의 지속성을 보여주는지 확인한다.','Citi':'유로존 직접 물가 전망이 충분히 확보되지 않았다. 국가 CPI를 근거로 확정적 ECB 반응을 만들어내지 않는다.','BofA':'광범위한 2차 효과 억제와 수요 부족이라는 관점을 시험한다. 헤드라인 에너지 상회만으로 12월 추가 인상을 기본 전망으로 바꾸지는 않는 해석이다.','HSBC':'국내 임금·2차 효과는 온건하지만 공급·지정학 위험은 상방이라는 관점이다. 국가 헤드라인과 지역 근원 압력의 차이를 우선 확인한다.'}},
      'GB':{
       'growth':{'GS':'물가가 인상을 정당화할 정도는 아니라는 동결 전망에 비춰 수요 압력을 점검한다. 특히 임금은 생산성·서비스 가격과 같이 봐야 한다.','JPM':'채용·임금의 초기 개선 신호와 노동시장 통계의 괴리를 함께 본다. 임금 압력의 지속성은 인상 조건과 관련되지만 당월 3개월 평균 임금을 2027년 임계값과 직접 비교하지 않는다.','Citi':'생산성 개선 가능성을 검토하는 관점이다. 임금 상회가 생산성으로 흡수되는지 확인해야 하며 월별 생산·고용만으로 AI의 인과효과를 판정하지 않는다.','BofA':'상반기 강세 이후 하반기 수요 둔화를 기본으로 본다. 에너지·금융여건 영향 아래 생산·소비가 계속 약하면 이 경로와 부합하며 상회가 넓어지면 재점검한다.','HSBC':'9월 동결과 에너지 2차 효과에 따른 11월 인상 위험을 구별한다. 성장 부진과 임금·서비스 물가가 함께 식는지 확인한다.'},
       'inflation':{'GS':'2026년 11월 헤드라인 정점 이후 2027년 빠른 둔화를 예상한다. 근원·서비스와 에너지 기저효과를 구분해 이 경로의 지속성을 확인한다.','JPM':'물가 초과의 높이보다 기간과 2027년 임금 지속성이 중요하다. 일시적 헤드라인 상승과 서비스 압력 고착을 분리한다.','Citi':'생산성의 물가 효과는 아직 확정하지 않았다. 이번 CPI가 좋아져도 생산성 개선이 원인이라고 단정하지 않는다.','BofA':'2026년 동결이 기본이지만 위험은 인상 방향이다. 근원·서비스의 지속적 상회가 인상 위험을 키우는 방향인지 확인한다.','HSBC':'에너지 충격의 2차 효과가 아직 나타나지 않았다는 전제다. 임금·서비스 가격으로 전가가 넓어지는지가 11월 위험을 가르는 조건이다.'},
       'policy':{'GS':'2026년 동결 후 2027년 인하라는 경로를 확인한다.','JPM':'7월 말 11월 인상 전망과 9월 물가 지속성 조건을 함께 보되, 오래된 경로라는 한계를 표시한다.','Citi':'현재 저장소에 확정적인 BoE 금리 경로가 없다. 생산성 분석을 금리 전망으로 바꾸지 않는다.','BofA':'2026년 동결·2027년 11월 한 차례 인하가 기본이며 인상은 위험 시나리오다.','HSBC':'9월 6대3의 3.75% 동결 예상과 실제 결정·표결을 대조하고, 11월 인상 위험 가이던스를 확인한다.'}},
      'CA':{
       'growth':{'GS':'고용·임금 약화가 2.25% 동결을 지지한다는 전제다. 약세가 지속되면 추가 인상 위험은 약해지는 방향이지만 저장된 기본 전망은 즉각 인하가 아니다.','JPM':'최근 고용 감소를 이전 반등의 모멘텀 약화로 본다. 전반적인 노동수요 악화인지, 일부 변동성인지 가려 남아 있는 유휴를 평가한다.','Citi':'약 1% 성장·초과공급이 기본이다. 재차 약한 지표는 이 관점과 부합하며, 지속적 회복이 확인되면 단순 변동성이라는 해석을 재검토해야 한다.','BofA':'무역 불확실성과 근원물가 안정에 따른 장기 동결을 기본으로 본다. 경기 약화만으로 동결을 인하로 바꾸지 않고 에너지 물가와 함께 판단한다.','HSBC':'무역에 따른 성장 하방과 물가 상방의 균형을 중시한다. 저장된 확정 금리 경로가 없어 특정 결정으로 치환하지 않는다.'},
       'inflation':{'GS':'근원물가·임금이 온건하다는 동결 전제와 대조한다. 헤드라인 에너지 상승이 근원으로 넓어지면 인상 위험을 재검토하는 구조다.','JPM':'초과공급·기업 가격결정력 제약이 전가를 막는다는 관점이다. 소비자가격뿐 아니라 생산자물가와 마진을 구분해 전가를 확인한다.','Citi':'수요 부진이 에너지 전가를 제약한다는 전제다. 낮은 물가는 이 관점과 부합하지만 BoC의 매파적 소통 때문에 즉각 인하로 연결하지 않는다.','BofA':'근원물가 안정이 2027년까지 동결의 전제다. 생산자·헤드라인 압력이 근원에 지속적으로 확산되는지를 본다.','HSBC':'성장 위험과 인플레이션 상방의 균형을 점검한다. 확정적인 다음 정책 경로는 기존 자료에서 확보하지 못했다.'}}
    }
    linked=macro.get(region,{}).get(theme,{}).get(house,'해당 지표가 아래의 수요·가격·정책 전제에 미치는 경로를 확인한다. 직접 수치 전망이 없어 방향을 단정하지 않는다.')
    if family in {'liquidity','fiscal'}:return '이 지표의 시장금리·유동성 경로를 설명하는 직접 IB 근거가 제한적이다. 정책금리 전망만으로 자금 유입·유출을 해석하지 않는다.'
    if future:return '실제값 미수록으로 강화·약화 판정은 보류한다. '+linked
    if family=='trade':signal='예상 대비와 수정 종전 대비 방향이 다를 수 있어 순수출 구성 확인 전에는 성장 전제의 강화·약화를 확정하지 않는다. '
    elif family in {'prices','wages','gdp','labor','production','pmi','housing','consumption','activity'}:
        surprise=e['surprise']
        signal='파일 컨센서스가 없어 시장 예상 대비 판정은 불가능하다. ' if surprise is None else '파일 컨센서스에 부합한 결과다. ' if abs(surprise)<1e-8 else ('파일 컨센서스보다 높은' if surprise>0 else '파일 컨센서스보다 낮은')+' 결과다. 시장 컨센서스 대비 평가이며 IB 자체 예상 대비와 구별한다. '
        if family=='labor' and any(t in e['event'].lower() for t in ['claims','unemployment','claimant']):signal+='높은 값은 노동시장 유휴 확대 방향일 수 있어 증가를 호재로 표시하지 않는다. '
    else:signal=''
    return signal+linked

COMMENT_SNAPSHOT_FIELDS=('survey','actual','prior','revised','scale','currency','unit')

def load_authored_comments(src,events,asof):
    """Load source-reviewed prose, never synthesize prose from FTS excerpts."""
    path=HERE/'authored_commentary.json'
    if not path.exists():return {},set()
    bundle=json.loads(path.read_text(encoding='utf-8'))
    if bundle.get('schema_version')!=1:raise ValueError('지원하지 않는 지표 코멘트 형식')
    current={e['id']:e for e in events};selected={};stale=set();seen=set()
    for review in sorted(bundle['event_reviews'],key=lambda r:r['reviewed_at']):
        if review['review_id'] in seen:raise ValueError('중복 코멘트 이력 ID')
        seen.add(review['review_id'])
        e=current.get(review['event_id'])
        if e is None or review['reviewed_on']>asof:continue
        if identity(review['match'])!=identity(e):raise ValueError('지표 코멘트 식별자 불일치')
        if review.get('review_status')!='draft' or review.get('origin')!='assistant_review_of_local_reports':raise ValueError('코멘트 작성 주체/검토 상태 오류')
        if any(not isinstance(review.get(k),str) or not review[k].strip() for k in ('headline','summary','watch','review_log')):raise ValueError('빈 지표 해석')
        if set(review['snapshot'])!=set(COMMENT_SNAPSHOT_FIELDS):raise ValueError('코멘트 숫자 스냅샷 누락')
        if any(review['snapshot'][k]!=e.get(k) for k in COMMENT_SNAPSHOT_FIELDS):
            stale.add(e['id']);selected.pop(e['id'],None);continue
        houses=[];seen_houses=set()
        for h in review['houses']:
            if h['house'] not in HOUSES or h['house'] in seen_houses:raise ValueError('코멘트 IB 중복/오류')
            seen_houses.add(h['house'])
            if h['kind'] not in {'reaction','preview','related'}:raise ValueError('코멘트 시점 분류 오류')
            if any(not h.get(k) for k in ('summary','comment','evidence','reviewed_pages')):raise ValueError('IB 요약/해석/근거 누락')
            if h['kind']=='related' and not h.get('limitation'):raise ValueError('관련 자료의 범위 제한 누락')
            refs=[evidence(src,x) for x in h['evidence']]
            read_pages={(x['doc_id'],x['page']) for x in h['reviewed_pages']}
            if any((x['doc_id'],x['page']) not in read_pages for x in refs):raise ValueError('전체 페이지 미검토 근거')
            if any(x['house']!=h['house'] or x['published']>asof for x in refs):raise ValueError('IB/근거 기준일 불일치')
            if h['kind']=='reaction' and (e['actual'] is None or any(x['published']<e['date'] for x in refs)):raise ValueError('발표 이전 원문을 사후 코멘트로 분류')
            if h['kind']=='preview' and any(x['published']>e['date'] for x in refs):raise ValueError('발표 이후 원문을 사전 전망으로 분류')
            houses.append({**h,'evidence':refs})
        selected[e['id']]={**review,'houses':houses}
    return selected,stale-set(selected)

def build(book,root,output,asof,include_interpretation=False,requested_at=None):
    events=load_calendar(book)
    calendar_events=events
    events=indicator_history.merge(calendar_events,indicator_history.read())
    src=sqlite3.connect((root/'data/research.sqlite3').as_uri()+'?mode=ro',uri=True);src.row_factory=sqlite3.Row
    state=sqlite3.connect((root/'data/living_ib.sqlite3').as_uri()+'?mode=ro',uri=True)
    views=[json.loads(r[0]) for r in state.execute('select payload from views')]
    forecasts=[json.loads(r[0]) for r in state.execute('select payload from forecasts')]
    review=json.loads((HERE/'reviewed_comments.json').read_text(encoding='utf-8'))
    views+=review.get('views',[])
    notes=review['comments'];note_map=collections.defaultdict(list)
    for n in notes:
        es=[evidence(src,e) for e in n['evidence']]
        if any(e['house']!=n['house'] for e in es):raise ValueError('IB 불일치')
        if any(e['published']>asof for e in es):continue
        note_map[(identity(n['match']),n['house'])].append({**n,'evidence':es})
    authored,stale_commentary=load_authored_comments(src,events,asof)
    source_refs={};counts=collections.Counter();used_notes=set();candidate_docs=set()
    for e in calendar_events:
        e['facts']=factual(e);e['reading']=interpretation(e) if include_interpretation else '';e['houses']=[]
        e['available_at_request']=e['actual'] is not None and e['date']<=asof
        e['commentary']=authored.get(e['id'])
        e['commentary_status']='source_reviewed_draft' if e['commentary'] else 'needs_recheck' if e['id'] in stale_commentary else 'not_reviewed'
        for house in HOUSES:
            direct=[dict(d) for d in note_map.get((identity(e),house),[])]
            for d in direct:
                if numeric(d.get('snapshot_actual'))!=e['actual']:
                    d['limitation']+=' 코멘트 검토 당시 파일 실제값과 현재값이 다릅니다. 기존 코멘트의 숫자 비교는 당시 기준이며 최신 값으로 재검토해야 합니다.'
            for d in direct:used_notes.add((identity(e),house,d['evidence'][0]['doc_id']))
            selected,region=select_views(views,e,house,asof)
            bases=[{k:v[k] for k in ['view_id','summary','effective_date','track','scope','assumptions','analyst_questions','limitations']}|{'context_only':v.get('context_only',False),'evidence':[evidence(src,x) for x in v['evidence']]} for v in selected]
            # An exact forecast is not a reaction and requires matching country, unit AND target period.
            fs=[f for f in forecasts if f['house']==house and f['country']==e['country_code'] and f['indicator']==e['indicator'] and f['reference_period']==e['reference_month'] and f['unit']==e['unit'] and f['as_of']<=asof and f['as_of']<e['date']]
            fg=collections.defaultdict(list)
            for f in fs:fg[f['scenario']].append(f)
            fs=[f for history in fg.values() for f in history if f['as_of']==max(x['as_of'] for x in history)]
            exact=[{**f,'difference':None if e['actual'] is None else round(e['actual']-f['value'],8),'evidence':[evidence(src,x) for x in f['evidence']]} for f in fs]
            candidates=indicator_candidates(src,e,house,asof)
            candidate_docs.update(x['doc_id'] for x in candidates)
            status='reaction' if any(d['kind']=='reaction' for d in direct) else 'preview' if direct or exact else 'candidate' if candidates else ('inferred' if include_interpretation else 'basis') if bases else 'uncovered'
            counts[status]+=1
            gap=('지표명·동의어와 국가/지역을 함께 검색한 미검토 원문 후보가 있다. 원문 전체를 확인하기 전에는 IB 코멘트로 인용하지 않는다.' if candidates else '지표명·동의어와 국가/지역을 함께 검색했지만 원문 후보를 확보하지 못했다.')
            analysis=stance_reading(e,house,bases,region) if include_interpretation else ''
            if bases:
                analysis=('유로존 공통 견해를 이 국가에 적용하는 추론이다. ' if region!=e['country_code'] else '')+analysis
            else:analysis='해당 국가의 적용 가능한 기존 IB 견해를 확보하지 못해 IB별 해석을 만들지 않았다. '+e['reading']
            entry={'house':house,'status':status,'direct':direct,'forecasts':exact,'search_candidates':candidates,'bases':bases,'analysis':analysis if include_interpretation else '', 'interpretation_status':'template_draft' if include_interpretation else 'not_requested','gap':gap,'regional_basis':region!=e['country_code'],'basis_after_release':any(v['effective_date']>e['date'] for v in selected)}
            entry['commentary']=next((h for h in (e['commentary'] or {}).get('houses',[]) if h['house']==house),None)
            if entry['commentary']:
                counts[status]-=1;entry['status']=entry['commentary']['kind'];counts[entry['status']]+=1
                entry['interpretation_status']='source_reviewed_draft'
                for ref in entry['commentary']['evidence']:source_refs[ref['doc_id']]=ref
            elif e['commentary_status']=='needs_recheck':
                counts[status]-=1;entry['status']='needs_recheck';counts[entry['status']]+=1
            e['houses'].append(entry)
            for item in direct+exact+bases:
                for ref in item['evidence']:
                    if not ref['published'] or ref['published']>asof:raise ValueError('기준일 이후/날짜 미확인 근거')
                    source_refs[ref['doc_id']]=ref
    for e in events:
        if e['record_origin']!='history':continue
        e['commentary']=authored.get(e['id'])
        e['commentary_status']='source_reviewed_draft' if e['commentary'] else 'needs_recheck' if e['id'] in stale_commentary else 'not_reviewed'
        e['available_at_request']=e['actual'] is not None and e['date']<=asof
        for h in e.get('houses',[]):
            h['commentary']=next((x for x in (e['commentary'] or {}).get('houses',[]) if x['house']==h['house']),None)
            if h['commentary']:
                h['status']=h['commentary']['kind']
                h['interpretation_status']='source_reviewed_draft'
                for x in h['commentary']['evidence']:source_refs[x['doc_id']]=x
            elif e['commentary_status']=='needs_recheck':
                h['status']='needs_recheck'
            for item in h.get('direct',[])+h.get('forecasts',[]):
                for x in item.get('evidence',[]):source_refs[x['doc_id']]=x
    counts=collections.Counter(h['status'] for e in events for h in e['houses'])
    dates=[e['date'] for e in events];months=sorted({e['release_month'] for e in events})
    columns=calendar_events[0].get('source_columns') or [{'letter':chr(66+i),'key':k,'label':LABELS[i]} for i,k in enumerate(FIELDS)]
    source_column_range=calendar_events[0].get('source_column_range','B:K')
    meta={'title':'매크로 지표와 IB 해석','as_of':asof,'generated_at':datetime.now().astimezone().isoformat(timespec='seconds'),'workbook':str(book),'workbook_name':book.name,'workbook_sha256':hashlib.sha256(book.read_bytes()).hexdigest(),'workbook_modified':datetime.fromtimestamp(book.stat().st_mtime).isoformat(timespec='seconds'),'source_db':str(root/'data/research.sqlite3'),'source_date_max':src.execute('select max(published) from documents where status="indexed"').fetchone()[0],'event_count':len(events),'country_count':len({e['country_code'] for e in events}),'reported':sum(e['actual'] is not None for e in events),'missing':sum(e['actual'] is None for e in events),'start':min(dates),'end':max(dates),'months':months,'comment_counts':dict(counts),'reviewed_connections':len(used_notes),'source_documents':len(source_refs),'indicator_candidate_documents':len(candidate_docs),'indicator_candidate_pages':sum(len(h['search_candidates']) for e in events for h in e['houses']),'time_note':'발표시각은 파일 D열 그대로이며 시간대는 파일에 명시되지 않음. 실제값 존재 여부는 H열 기준. 원문 후보는 지표 키워드와 국가/지역을 결합한 미검토 검색 결과이며, 직접 IB 코멘트로 사용하기 전에 전체 페이지 검토가 필요함. 상위 거시 뷰는 보조 배경일 뿐 지표 코멘트가 아님.','monthly_note':'월 분류는 발표월 기준. 기간 밖으로 밀려난 실제값 수록 지표는 이전에 저장한 원본 수치와 검토 당시 코멘트로 누적 표시함. 최신 원본에 같은 지표가 다시 있으면 최신 수치를 우선함.'}
    meta.update({'source_column_range':source_column_range,'requested_at':requested_at or datetime.now().astimezone().isoformat(timespec='seconds'),'interpretation_mode':'reviewed_template' if include_interpretation else 'data_and_existing_evidence','available_at_request':sum(e['available_at_request'] for e in events)})
    meta.update(latest_calendar_start=min(e['date'] for e in calendar_events),latest_calendar_end=max(e['date'] for e in calendar_events),latest_calendar_event_count=len(calendar_events),latest_calendar_reported=sum(e['actual'] is not None for e in calendar_events),history_event_count=len(events)-len(calendar_events))
    meta.update(authored_event_count=len(authored),authored_house_count=sum(len(r['houses']) for r in authored.values()),commentary_needs_recheck=len(stale_commentary),commentary_reviewed_on=max((r['reviewed_on'] for r in authored.values()),default=None))
    if authored:meta['interpretation_mode']='source_reviewed_commentary'
    if source_column_range=='A:L':
        meta['time_note']='발표일시는 파일 A열 그대로이며 시간대는 파일에 명시되지 않음. 날짜만 있는 행은 시각 미확인으로 보존. 실제값은 H열 기준이며 % 표시형식의 분수는 퍼센트 수치로 변환. k/m/b/t와 통화기호는 원본에 명시된 값만 사용. 원문 후보는 전체 페이지 검토가 필요한 미검토 검색 결과이고 상위 거시 뷰는 보조 배경임.'
    macro_spec=importlib.util.spec_from_file_location('daily_macro_context',HERE/'build_macro_views.py')
    macro=importlib.util.module_from_spec(macro_spec);macro_spec.loader.exec_module(macro)
    meta['macro_views']=macro.build(root,output,asof,requested_at)
    payload={'schema_version':2,'meta':meta,'columns':columns,'countries':[{'code':c,'label':l} for c,l in COUNTRIES.values() if any(e['country_code']==c for e in events)],'themes':THEMES,'houses':HOUSES,'events':events,'sources':list(source_refs.values())}
    output.mkdir(parents=True,exist_ok=True)
    encoded=json.dumps(payload,ensure_ascii=False,allow_nan=False,separators=(',',':')).replace('</','<\\/')
    (output/'dashboard_data.json').write_text(encoded,encoding='utf-8')
    template=(HERE/'dashboard.html').read_text(encoding='utf-8')
    (output/'index.html').write_text(template.replace('__DASHBOARD_DATA__',encoded),encoding='utf-8')
    (output/'validation.json').write_text(json.dumps({'rows_preserved':len(events),'unique_events':len({e['id'] for e in events}),'columns_preserved':source_column_range,'five_houses_per_event':all(len(e['houses'])==5 for e in events),'quote_and_file_checks':'passed','status_counts':dict(counts),'direct_event_house_connections':len(used_notes),'indicator_candidate_documents':len(candidate_docs),'indicator_candidate_pages':meta['indicator_candidate_pages'],'indicator_search_is_unreviewed':True},ensure_ascii=False,indent=2),encoding='utf-8')
    src.close();state.close();return meta

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--workbook',type=Path);p.add_argument('--root',type=Path,default=DEFAULT_ROOT);p.add_argument('--output',type=Path,default=HERE/'dist');p.add_argument('--as-of',default=date.today().isoformat());p.add_argument('--include-reviewed-interpretations',action='store_true',help='이전에 검토한 고정 해석 템플릿을 표시. 새 AI 분석을 실행하는 옵션이 아닙니다.');a=p.parse_args()
    candidates=sorted(DEFAULT_RAW.glob('ecocal*.xlsx'),key=lambda p:(p.stat().st_mtime,p.name),reverse=True)
    if not a.workbook and not candidates:raise SystemExit('ecocal*.xlsx 파일 없음')
    print(json.dumps(build(a.workbook or candidates[0],a.root,a.output,a.as_of,a.include_reviewed_interpretations),ensure_ascii=False,indent=2))
